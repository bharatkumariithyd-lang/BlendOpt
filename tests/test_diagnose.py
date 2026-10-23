"""Stage 6 tests: diagnosis of infeasible problems (pre-checks and the elastic model)."""

from dataclasses import replace

import openpyxl
import pytest

from blend_optimizer import config, examples
from blend_optimizer import solve_and_report as sr
from blend_optimizer.diagnose import best_mix, best_value_alone, conflict_set, precheck, usage_bounds
from blend_optimizer.inputs import load_inputs
from blend_optimizer.linear_rows import build_linear_rows, violation_scale
from blend_optimizer.model import build_model
from blend_optimizer.pipeline import best_achievable, diagnose_infeasible, optimise


def _load(write_example, example, file_name="product.xlsx"):
    return load_inputs(write_example(example, file_name))


def _with_spec(example, name, **cells):
    """Copy of ``example`` with cells of one Specs row changed, e.g. Min=12."""
    specs = example.specs.copy()
    for column, value in cells.items():
        specs.loc[specs[config.COL_NAME] == name, column] = value
    return replace(example, specs=specs)


def _with_raw(example, name, **cells):
    """Copy of ``example`` with cells of one RawMaterials row changed."""
    table = example.raw_materials.copy()
    for column, value in cells.items():
        table.loc[table[config.COL_NAME] == name, column] = value
    return replace(example, raw_materials=table)


def _text(diagnosis):
    return "\n".join(diagnosis.messages)


# ===========================================================================
# Acceptance test 1: the fluorspar ban
# ===========================================================================

def test_fluorspar_ban_is_caught_by_the_precheck(write_example):
    data = _load(write_example, examples.synthetic_slag_example(ban_fluorspar=True))
    findings = precheck(data)
    assert len(findings) == 1
    finding = findings[0]
    assert finding.row.name == "CaF2_min"
    assert "No allowed raw material contains CaF2, so CaF2_min" in finding.message
    assert "cannot be met" in finding.message
    assert "Fluorspar (CaF2 85.00 %) is not allowed (Max 0 %, UsageLimits)" in finding.message
    assert finding.best == 0

    diagnosis = diagnose_infeasible(data)
    assert finding.message in _text(diagnosis)
    # The elastic model agrees: CaF2_min and the fluorspar ban are what conflict
    assert diagnosis.conflict.spec_rows == ["CaF2_min"]
    assert diagnosis.conflict.usage_limits == [("Fluorspar", "max", 0.0)]
    assert "lower CaF2_min from 2.00 % to 0.00 %" in _text(diagnosis)


# ===========================================================================
# Acceptance test 2: Test 1 with Al2O3 Min 45
# ===========================================================================

def test_al2o3_45_is_a_conflict_between_cao_and_al2o3(write_example):
    data = _load(write_example, examples.hand_worked_example(al2o3_min=45))
    assert precheck(data) == []  # each limit can be met on its own

    diagnosis = diagnose_infeasible(data)
    assert diagnosis.conflict.spec_rows == ["CaO_min", "Al2O3_min"]
    text = _text(diagnosis)
    assert "Pre-check: every limit can be met on its own" in text
    assert "these limits cannot all be met together: CaO_min, Al2O3_min" in text

    # Max Al2O3 with CaO >= 50: lime + alumina with CaO = 50 -> x_lime = 49/91,
    # Al2O3 = (49 * 1 + 42 * 85) / 91 = 39.77 %
    assert ("Al2O3_min (Al2O3 >= 45.00 %): the highest Al2O3 achievable while CaO >= 50.00 % "
            "and all other limits are held is 39.77 %.") in text
    # The other way round: max CaO with Al2O3 >= 45: x_lime = 40/84, CaO = (40 * 92 + 44 * 1) / 84 = 44.33 %
    assert ("CaO_min (CaO >= 50.00 %): the highest CaO achievable while Al2O3 >= 45.00 % "
            "and all other limits are held is 44.33 %.") in text

    table = diagnosis.table.set_index("Limit")
    assert table.loc["Al2O3_min", "Best achievable"] == pytest.approx(3619 / 91, abs=1e-4)
    assert table.loc["CaO_min", "Best achievable"] == pytest.approx(3724 / 84, abs=1e-4)
    assert table.loc["Al2O3_min", "Change needed"] == pytest.approx(45 - 3619 / 91, abs=1e-4)


def test_al2o3_45_elastic_suggestion(write_example):
    data = _load(write_example, examples.hand_worked_example(al2o3_min=45))
    elastic = sr.solve(build_model(data, elastic=True))
    # Relaxing CaO to 44.33 is 11.3 % of its limit; relaxing Al2O3 to 39.77 would be 11.6 %
    assert elastic.violations["CaO_min"] == pytest.approx(50 - 3724 / 84, abs=1e-6)
    assert elastic.violations["Al2O3_min"] == 0
    assert elastic.relative_violation == pytest.approx((50 - 3724 / 84) / 50)
    assert "lower CaO_min from 50.00 % to 44.33 %" in _text(diagnose_infeasible(data))


# ===========================================================================
# More cases
# ===========================================================================

def test_ratio_conflict_matches_stage3_finding(write_example):
    # Stage 3 found the lowest achievable basicity with the other Test 2 limits held: 11.2378
    data = _load(write_example, _with_spec(examples.synthetic_slag_example(), "Basicity", Max=11))
    assert precheck(data) == []
    diagnosis = diagnose_infeasible(data)
    assert "Basicity_max" in diagnosis.conflict.spec_rows
    best = diagnosis.table.set_index("Limit").loc["Basicity_max", "Best achievable"]
    assert best == pytest.approx(11.2378, abs=1e-3)
    assert "the lowest CaO / SiO2 achievable while" in _text(diagnosis)


def test_usage_minimums_above_100(write_example):
    limits = examples.usage_limit_table([("Lime", 60, None), ("Raw bauxite", 50, None)])
    data = _load(write_example, replace(examples.synthetic_slag_example(), usage_limits=limits))
    findings = precheck(data)
    assert [f.row for f in findings] == [None]
    assert "The minimum usage limits add up to 110.00 % of the blend, more than 100 %" in findings[0].message
    assert "Lime min 60 %" in findings[0].message
    diagnosis = diagnose_infeasible(data)
    assert "Elastic model: no answer" in _text(diagnosis)


def test_usage_maximums_below_100(write_example):
    example = examples.hand_worked_example()
    for name in ("Lime", "Calcined alumina", "Bauxite"):
        example = _with_raw(example, name, **{config.COL_MAX_USE: 30})
    findings = precheck(_load(write_example, example))
    assert "The maximum usage limits add up to only 90.00 %" in findings[0].message


def test_precheck_max_row(write_example):
    data = _load(write_example, _with_spec(examples.hand_worked_example(), "SiO2", Max=1))
    findings = precheck(data)
    assert len(findings) == 1
    assert ("SiO2_max (SiO2 <= 1.00 %) cannot be met by any mix of the allowed raw materials, even ignoring "
            "all other spec limits: the lowest achievable SiO2 is 2.00 %.") in findings[0].message


def test_precheck_capped_materials_are_named(write_example):
    example = examples.hand_worked_example()
    example = _with_raw(example, "Calcined alumina", **{config.COL_MAX_USE: 10})
    example = _with_raw(example, "Bauxite", **{config.COL_MAX_USE: 10})
    findings = precheck(_load(write_example, example))
    message = findings[0].message
    # Best possible: 10 % alumina + 10 % bauxite + 80 % lime = 8.5 + 6.0 + 0.8 = 15.30 % Al2O3
    assert "the highest achievable Al2O3 is 15.30 %" in message
    assert "Calcined alumina (Al2O3 85.00 %) is capped at 10 % (plant-wide limit in RawMaterials)" in message
    assert "Bauxite (Al2O3 60.00 %) is capped at 10 %" in message


def test_precheck_ratio_row(write_example):
    data = _load(write_example, _with_spec(examples.synthetic_slag_example(), "Basicity", Min=100))
    findings = precheck(data)
    assert [f.row.name for f in findings] == ["Basicity_min"]
    # Best single material: lime, 92 / 1.5 = 61.33
    assert "the highest achievable CaO / SiO2 is 61.33." in findings[0].message


def test_precheck_on_ignited_basis(write_example):
    example = _with_spec(examples.synthetic_slag_example(spec_basis="ignited"), "Al2O3", Min=90, Max=95)
    findings = precheck(_load(write_example, example))
    # Calcined bauxite alone: 85 / (1 - 1.0/100) = 85.86 % Al2O3 on an ignited basis
    assert "the highest achievable Al2O3 is 85.86 %" in findings[0].message
    assert findings[0].best == pytest.approx(85 / 0.99, abs=1e-6)


def test_feasible_problem_has_nothing_to_report(write_example):
    data = _load(write_example, examples.synthetic_slag_example())
    assert precheck(data) == []
    elastic = sr.solve(build_model(data, elastic=True))
    assert elastic.relative_violation == pytest.approx(0, abs=1e-9)
    assert conflict_set(data, elastic).spec_rows == []


def test_conflict_where_the_others_also_conflict(write_example):
    # Al2O3 45 conflicts with CaO 50, and SiO2 <= 1 is impossible on its own
    example = _with_spec(examples.hand_worked_example(al2o3_min=45), "SiO2", Max=1)
    data = _load(write_example, example)
    assert [f.row.name for f in precheck(data)] == ["SiO2_max"]
    cao = next(row for row in build_linear_rows(data) if row.name == "CaO_min")
    assert best_achievable(data, cao) is None  # even without CaO_min, SiO2 <= 1 cannot be met


# ===========================================================================
# Building blocks
# ===========================================================================

def test_best_mix_fills_the_best_material_first():
    bounds = {"A": (0.1, 0.5), "B": (0.0, 1.0), "C": (0.2, 0.3)}
    coefficients = {"A": 10.0, "B": 1.0, "C": 5.0}
    # highest: A to 0.5, then C to 0.3, then B the remaining 0.2
    assert best_mix(coefficients, bounds, highest=True) == pytest.approx(10 * 0.5 + 5 * 0.3 + 1 * 0.2)
    # lowest: B takes everything above the minimums (A 0.1, C 0.2)
    assert best_mix(coefficients, bounds, highest=False) == pytest.approx(10 * 0.1 + 5 * 0.2 + 1 * 0.7)
    assert best_mix(coefficients, {"A": (0.6, 1.0), "B": (0.6, 1.0)}) is None   # minimums > 100 %


def test_best_value_alone_ratio(write_example):
    data = _load(write_example, examples.synthetic_slag_example())
    row = next(row for row in build_linear_rows(data) if row.name == "Basicity_min")
    bounds = usage_bounds(data)
    assert best_value_alone(row, bounds, highest=True) == pytest.approx(92 / 1.5)
    assert best_value_alone(row, bounds, highest=False) == pytest.approx(1.5 / 9.0)  # fluorspar alone


def test_violation_scale(write_example):
    data = _load(write_example, examples.synthetic_slag_example())
    rows = {row.name: row for row in build_linear_rows(data)}
    assert violation_scale(rows["CaO_min"], data) == 50            # dry row: the limit itself
    sio2 = [1.5, 2.5, 1.5, 6.0, 9.0, 9.0]                           # SiO2 of the six materials
    assert violation_scale(rows["Basicity_min"], data) == pytest.approx(8 * sum(sio2) / 6)
    zero = replace(rows["P_max"], limit=0.0)
    assert violation_scale(zero, data) == config.ELASTIC_MIN_LIMIT  # never divides by zero


def test_elastic_model_lp_file(write_example, tmp_path):
    data = _load(write_example, examples.hand_worked_example(al2o3_min=45))
    model = build_model(data, elastic=True)
    assert set(model.violations) == {"CaO_min", "Al2O3_min", "SiO2_max"}
    result = sr.solve(model, lp_path=tmp_path / "elastic.lp")
    text = (tmp_path / "elastic.lp").read_text()
    assert "CaO_min: viol_CaO_min + 2 x_Bauxite + x_Calcined_alumina + 92 x_Lime >= 50" in text
    assert "SiO2_max: - viol_SiO2_max" in text
    assert result.elastic and result.is_optimal


# ===========================================================================
# In the report
# ===========================================================================

def test_diagnosis_runs_automatically_in_the_report(write_example, tmp_path, capsys):
    data = _load(write_example, examples.synthetic_slag_example(ban_fluorspar=True))
    outcome = optimise(data, tmp_path / "ban.xlsx")
    assert outcome.diagnosis is not None
    assert (tmp_path / "ban_elastic.lp").is_file()

    printed = capsys.readouterr().out
    assert "No allowed raw material contains CaF2, so CaF2_min" in printed

    workbook = openpyxl.load_workbook(outcome.excel_path)
    assert workbook.sheetnames == ["Summary", "Diagnosis"]
    sheet = workbook["Diagnosis"]
    cells = [cell.value for cell in sheet["A"] if cell.value]
    assert "CaF2_min" in cells and "Fluorspar max use" in cells
    assert any("No allowed raw material contains CaF2" in str(cell) for cell in cells)


def test_feasible_run_has_no_diagnosis(write_example):
    outcome = optimise(_load(write_example, examples.hand_worked_example()), show=False)
    assert outcome.diagnosis is None and "Diagnosis" not in outcome.tables


def test_diagnosis_table_display(write_example):
    data = _load(write_example, examples.hand_worked_example(al2o3_min=45))
    shown = sr.display_table(diagnose_infeasible(data).table).set_index("Limit")
    assert shown.loc["Al2O3_min", "Best achievable"] == "39.77"
    assert shown.loc["Al2O3_min", "Elastic model target"] == ""   # not the limit the elastic model moved
    assert shown.loc["CaO_min", "Elastic model target"] == "44.33"

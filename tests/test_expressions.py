"""Conversion factors in spec expressions, e.g. Al2O3+1.8895*Al (example data only).

Hand calculations use two made-up raw materials:
    P: Al2O3 60 %, Al 10 %  ->  Al2O3+1.8895*Al = 60 + 18.895   =  78.895
    Q: Al2O3 72 %, Al 19 %  ->  Al2O3+1.8895*Al = 72 + 35.9005  = 107.9005
"""

from dataclasses import replace

import openpyxl
import pytest

from blend_optimizer import config, examples
from blend_optimizer import solve_and_report as sr
from blend_optimizer.data import InputError, Term
from blend_optimizer.diagnose import precheck
from blend_optimizer.examples import ExampleWorkbook
from blend_optimizer.inputs import load_inputs, parse_expression
from blend_optimizer.linear_rows import achieved_value, build_linear_rows
from blend_optimizer.model import build_model
from blend_optimizer.pipeline import diagnose_infeasible
from blend_optimizer.what_if import run_scenarios, what_if

FACTOR = 1.8895
TOTAL = "Al2O3+1.8895*Al"


# ===========================================================================
# Example data
# ===========================================================================

def pq_example(specs, p_use=(None, None), q_use=(None, None), basis="dry", loi=(0, 0), extra_column=False):
    """Raw materials P and Q (P is cheaper) with the given Specs rows.

    Args:
        specs: Spec rows: (Name, Numerator, Denominator, Min, Max, Margin).
        p_use, q_use: (Min_use_%, Max_use_%) for P and Q; (65, 65) fixes P at 65 %.
        basis: Spec basis.
        loi: LOI_% of P and Q (Others is reduced so each analysis still totals 100).
        extra_column: Add the old helper column Al_as_Al2O3 = 1.8895 * Al.
    """
    components = ["Al2O3", "Al", "Others"] + (["Al_as_Al2O3"] if extra_column else [])
    p = ["P", 10000, "dry", 0, loi[0], p_use[0], p_use[1], 60, 10, 30 - loi[0]]
    q = ["Q", 15000, "dry", 0, loi[1], q_use[0], q_use[1], 72, 19, 9 - loi[1]]
    if extra_column:
        p.append(FACTOR * 10)
        q.append(FACTOR * 19)
    return ExampleWorkbook(
        raw_materials=examples.raw_material_table(components, [p, q]),
        specs=examples.spec_table(specs),
        settings=examples.settings_dict("Factor test", basis),
    )


def _solve(write_example, example, file_name="product.xlsx"):
    data = load_inputs(write_example(example, file_name))
    return data, sr.solve(build_model(data))


def _check(data, result):
    return sr.spec_check_table(data, result).set_index("Spec")


# ===========================================================================
# The parser
# ===========================================================================

@pytest.mark.parametrize("text, expected, shown", [
    ("Al2O3", [Term(1.0, "Al2O3")], "Al2O3"),
    ("Al2O3+Al", [Term(1.0, "Al2O3"), Term(1.0, "Al")], "Al2O3+Al"),
    ("Al2O3+1.8895*Al", [Term(1.0, "Al2O3"), Term(1.8895, "Al")], "Al2O3+1.8895*Al"),
    (" Al2O3 + 1.8895 * Al ", [Term(1.0, "Al2O3"), Term(1.8895, "Al")], "Al2O3+1.8895*Al"),
    ("2*Si", [Term(2.0, "Si")], "2*Si"),
    (".5*Si", [Term(0.5, "Si")], ".5*Si"),
    ("CaO+1.3992*Ca+0.8748*CaC2", [Term(1.0, "CaO"), Term(1.3992, "Ca"), Term(0.8748, "CaC2")],
     "CaO+1.3992*Ca+0.8748*CaC2"),
    ("Al2O3+1*Al", [Term(1.0, "Al2O3"), Term(1.0, "Al")], "Al2O3+1*Al"),   # explicit 1 is kept as typed
    ("Fe-total", [Term(1.0, "Fe-total")], "Fe-total"),                     # hyphen inside a name
])
def test_valid_expressions(text, expected, shown):
    terms = parse_expression(text)
    assert terms == expected
    assert config.COMPONENT_SEPARATOR.join(term.label for term in terms) == shown


@pytest.mark.parametrize("text, message", [
    ("1.8895Al", "starts with a number; put '*' between the factor and the component (did you mean 1.8895*Al?)"),
    ("*Al", "has no factor before '*'"),
    ("1.8895*", "has no component after '*'"),
    ("Al*1.8895", "has the factor after the component; write it first: 1.8895*Al"),
    ("Al2O3++Al", "has an empty component name"),
    ("Al2O3+", "has an empty component name"),
    ("x*Al", "the factor 'x' is not a number"),
    ("1,8895*Al", "the factor '1,8895' is not a number (use a decimal point, e.g. 1.8895)"),
    ("0*Al", "the factor of Al is 0; it must be greater than 0"),
    ("-1.8895*Al", "subtraction ('-') is not supported"),
    ("Al2O3 - Al", "subtraction ('-') is not supported"),
    ("Al2O3-1.8895*Al", "subtraction ('-') is not supported"),
    ("2*3*Al", "has more than one '*'"),
    ("Al+1.8895*Al", "lists Al more than once"),
    ("", "the expression is empty"),
])
def test_invalid_expressions(text, message):
    with pytest.raises(ValueError) as info:
        parse_expression(text)
    assert message in str(info.value)


def test_component_names_take_priority():
    known = ["CaO", "MgO", "3CaO", "Fe-total"]
    assert parse_expression("3CaO", known) == [Term(1.0, "3CaO")]          # a real column wins
    assert parse_expression("Fe-total", known) == [Term(1.0, "Fe-total")]
    with pytest.raises(ValueError, match=r"did you mean 2\*CaO\?"):
        parse_expression("2CaO", known)                                   # no such column
    with pytest.raises(ValueError, match="subtraction"):
        parse_expression("CaO-MgO", known)                                # both parts are columns


def test_messages_name_sheet_row_and_expression(write_example):
    example = pq_example([
        ("Total Al2O3", "Al2O3+1.8895Al", None, 75, None, None),
        ("Typo", "Al2O3+1.8895*Alu", None, 75, None, None),
        ("Ratio", "Al2O3", "Others+0*Al", 1, None, None),
    ])
    with pytest.raises(InputError) as info:
        load_inputs(write_example(example))
    errors = info.value.errors
    assert errors[0].startswith("Specs (product.xlsx) row 2 'Total Al2O3', Numerator: in 'Al2O3+1.8895Al', "
                                "the term '1.8895Al' starts with a number")
    assert ("Specs (product.xlsx) row 3 'Typo', Numerator: component 'Alu' is not a column in RawMaterials. "
            "Did you mean 'Al'?") in errors[1]
    assert errors[2].startswith("Specs (product.xlsx) row 4 'Ratio', Denominator: in 'Others+0*Al'")


def test_factor_expression_may_exceed_100(write_example):
    # Al2O3+1.8895*Al can exceed 100 % (Q is 107.90), so Min 105 is not rejected as "never met"
    data = load_inputs(write_example(pq_example([("Total Al2O3", TOTAL, None, 105, None, None)])))
    assert data.specs[0].min_value == 105
    with pytest.raises(InputError, match="above 100 % and can never be met"):  # still applies to plain sums
        load_inputs(write_example(pq_example([("Al2O3", "Al2O3", None, 101, None, None)]), "plain.xlsx"))


# ===========================================================================
# Hand-calculated blend: P 65 %, Q 35 %
# ===========================================================================

def test_hand_calculated_blend(write_example):
    example = pq_example([("Al2O3", "Al2O3", None, 50, None, None),
                          ("Al", "Al", None, None, 20, None),
                          ("Total Al2O3", TOTAL, None, 80, None, None)],
                         p_use=(65, 65), q_use=(35, 35))
    data, result = _solve(write_example, example)
    assert result.fractions == pytest.approx({"P": 0.65, "Q": 0.35})
    check = _check(data, result)
    # Al2O3 = 0.65*60 + 0.35*72;  Al = 0.65*10 + 0.35*19;  total = Al2O3 + 1.8895 * Al
    assert check.loc["Al2O3", "Achieved"] == pytest.approx(64.20)
    assert check.loc["Al", "Achieved"] == pytest.approx(13.15)
    assert check.loc["Total Al2O3", "Achieved"] == pytest.approx(64.20 + FACTOR * 13.15)
    assert round(check.loc["Total Al2O3", "Achieved"], 2) == 89.05
    shown = sr.display_table(sr.spec_check_table(data, result)).set_index("Spec")
    assert shown.loc["Total Al2O3", "Limits"] == "Al2O3+1.8895*Al"
    assert shown.loc["Total Al2O3", "Achieved"] == "89.05"


# ===========================================================================
# End to end: Min 75
# ===========================================================================

def test_min_75_infeasible_plain_feasible_with_factor(write_example):
    plain = pq_example([("Al2O3", "Al2O3", None, 75, None, None)])
    data, result = _solve(write_example, plain, "plain.xlsx")
    assert result.status == "Infeasible"  # neither material has more than 72 % Al2O3
    assert "the highest achievable Al2O3 is 72.00 %" in precheck(data)[0].message

    factored = pq_example([("Total Al2O3", TOTAL, None, 75, None, None)])
    data, result = _solve(write_example, factored, "factored.xlsx")
    assert result.is_optimal
    row = next(row for row in build_linear_rows(data) if row.name == "Total Al2O3_min")
    # P = 60 + 1.8895*10 = 78.895 (78.90 to 2 dp), Q = 72 + 1.8895*19 = 107.9005 (107.90)
    assert row.numerator == pytest.approx({"P": 78.895, "Q": 107.9005})
    # P alone (78.90) already meets 75 and is cheaper
    assert result.fractions["P"] == pytest.approx(1.0)
    assert result.cost_per_dry_t == pytest.approx(10000)


def test_binding_factor_spec_and_shadow_price(write_example):
    data, result = _solve(write_example, pq_example([("Total Al2O3", TOTAL, None, 80, None, None)]))
    # 78.895 + (107.9005 - 78.895) x_Q = 80  ->  x_Q = 1.105 / 29.0055
    x_q = 1.105 / 29.0055
    assert result.fractions["Q"] == pytest.approx(x_q)
    assert result.cost_per_dry_t == pytest.approx(10000 + 5000 * x_q)
    shadow = sr.shadow_price_table(data, result).set_index("Constraint")
    assert shadow.loc["Total Al2O3_min", "Limits"] == "Al2O3+1.8895*Al"
    # One more % of the total costs (15000 - 10000) / 29.0055 Rs/t
    assert shadow.loc["Total Al2O3_min", "Saving per 1 unit (Rs/dry t)"] == pytest.approx(5000 / 29.0055)


# ===========================================================================
# Equivalence with the old ways
# ===========================================================================

def test_factor_one_equals_plain_sum(write_example):
    with_one = pq_example([("Total", "Al2O3+1*Al", None, 75, None, None)])
    plain = pq_example([("Total", "Al2O3+Al", None, 75, None, None)])
    data_one, result_one = _solve(write_example, with_one, "one.xlsx")
    data_plain, result_plain = _solve(write_example, plain, "plain.xlsx")
    assert data_one.specs[0].numerator == data_plain.specs[0].numerator
    assert result_one.fractions == result_plain.fractions
    assert result_one.cost_per_dry_t == result_plain.cost_per_dry_t
    assert result_one.row_dual == result_plain.row_dual
    assert _check(data_one, result_one).loc["Total", "Limits"] == "Al2O3+1*Al"   # shown as typed
    assert _check(data_plain, result_plain).loc["Total", "Limits"] == "Al2O3+Al"


def test_old_helper_column_gives_the_same_answer(write_example):
    new = pq_example([("Total", TOTAL, None, 80, None, None)])
    old = pq_example([("Total", "Al2O3+Al_as_Al2O3", None, 80, None, None)], extra_column=True)
    _, result_new = _solve(write_example, new, "new.xlsx")
    _, result_old = _solve(write_example, old, "old.xlsx")
    assert result_new.fractions == pytest.approx(result_old.fractions, abs=1e-9)
    assert result_new.cost_per_dry_t == pytest.approx(result_old.cost_per_dry_t, abs=1e-6)
    assert result_new.row_dual["Total_min"] == pytest.approx(result_old.row_dual["Total_min"])


# ===========================================================================
# A factor in a ratio's Denominator
# ===========================================================================
# A: CaO 90, SiO2 2, Si 0  ->  denominator SiO2+2.1393*Si = 2
# B: CaO 40, SiO2 10, Si 5 ->  denominator = 10 + 2.1393 * 5 = 20.6965

def ab_example(r_min, a_use=(None, None), b_use=(None, None)):
    rows = [
        ["A", 12000, "dry", 0, 0, a_use[0], a_use[1], 90, 2, 0, 8],
        ["B", 8000, "dry", 0, 0, b_use[0], b_use[1], 40, 10, 5, 45],
    ]
    return ExampleWorkbook(
        raw_materials=examples.raw_material_table(["CaO", "SiO2", "Si", "Others"], rows),
        specs=examples.spec_table([("Basicity", "CaO", "SiO2+2.1393*Si", r_min, None, None)]),
        settings=examples.settings_dict("Ratio factor test"),
    )


def test_ratio_with_factor_in_denominator_fixed_blend(write_example):
    data, result = _solve(write_example, ab_example(1, a_use=(50, 50), b_use=(50, 50)))
    # CaO = 0.5*90 + 0.5*40 = 65;  denominator = 0.5*2 + 0.5*20.6965 = 11.34825
    check = _check(data, result)
    assert check.loc["Basicity", "Achieved"] == pytest.approx(65 / 11.34825)
    assert check.loc["Basicity", "Limits"] == "CaO / (SiO2+2.1393*Si)"


def test_ratio_with_factor_in_denominator_binding(write_example):
    data, result = _solve(write_example, ab_example(5))
    # B is cheaper; least A for CaO / D >= 5:
    # (40 + 50 x) = 5 (20.6965 - 18.6965 x)  ->  x = (5*20.6965 - 40) / (50 + 5*18.6965)
    x_a = (5 * 20.6965 - 40) / (50 + 5 * 18.6965)
    assert result.fractions["A"] == pytest.approx(x_a)
    assert _check(data, result).loc["Basicity", "Achieved"] == pytest.approx(5)


# ===========================================================================
# Ignited basis with a factor
# ===========================================================================
# P LOI 5 %, Q LOI 2 %: 1 dry t leaves 0.95 t (P) and 0.98 t (Q) after ignition

def test_ignited_basis_with_factor_fixed_blend(write_example):
    example = pq_example([("Total Al2O3", TOTAL, None, 80, None, None)], p_use=(65, 65), q_use=(35, 35),
                         basis="ignited", loi=(5, 2))
    data, result = _solve(write_example, example)
    dry_total = 0.65 * 78.895 + 0.35 * 107.9005     # 89.046925 on a dry basis
    ignited_mass = 0.65 * 0.95 + 0.35 * 0.98        # 0.9605 t per dry t
    assert _check(data, result).loc["Total Al2O3", "Achieved"] == pytest.approx(dry_total / ignited_mass)
    assert round(dry_total / ignited_mass, 2) == 92.71


def test_ignited_basis_with_factor_binding(write_example):
    example = pq_example([("Total Al2O3", TOTAL, None, 90, None, None)], basis="ignited", loi=(5, 2))
    data, result = _solve(write_example, example)
    # (78.895 + 29.0055 x) = 90 (0.95 + 0.03 x)  ->  x_Q = (85.5 - 78.895) / (29.0055 - 2.7)
    x_q = (85.5 - 78.895) / (29.0055 - 2.7)
    assert result.fractions["Q"] == pytest.approx(x_q)
    assert _check(data, result).loc["Total Al2O3", "Achieved"] == pytest.approx(90)


# ===========================================================================
# Margins, what-if, diagnosis and the report
# ===========================================================================

def test_margin_on_a_factor_spec(write_example):
    data, result = _solve(write_example, pq_example([("Total Al2O3", TOTAL, None, 79.5, None, 0.5)]))
    check = _check(data, result)
    assert check.loc["Total Al2O3", "Target Min"] == 80
    assert check.loc["Total Al2O3", "Achieved"] == pytest.approx(80)


def test_what_if_on_a_factor_spec(write_example):
    data = load_inputs(write_example(pq_example([("Total Al2O3", TOTAL, None, 75, None, None)])))
    tables = what_if(data, spec_changes={"Total Al2O3": {"Min": 80}})
    cost = tables["Cost"].set_index("Item")
    assert cost.loc["Cost (Rs per dry tonne of blend)", "Scenario"] == pytest.approx(10000 + 5000 * 1.105 / 29.0055)
    assert tables["Spec check"].set_index("Spec").loc["Total Al2O3", "Limits"] == "Al2O3+1.8895*Al"

    # The same change from a WhatIf sheet
    example = replace(pq_example([("Total Al2O3", TOTAL, None, 75, None, None)]),
                      what_if=examples.what_if_table([("Higher total", "Spec Min", "Total Al2O3", 80)]))
    sheet_data = load_inputs(write_example(example, "whatif.xlsx"))
    outcome = run_scenarios(sheet_data)[0]
    assert outcome.result.cost_per_dry_t == pytest.approx(10000 + 5000 * 1.105 / 29.0055)


def test_diagnosis_shows_the_expression(write_example):
    data = load_inputs(write_example(pq_example([("Total Al2O3", TOTAL, None, 110, None, None)])))
    diagnosis = diagnose_infeasible(data)
    text = "\n".join(diagnosis.messages)
    assert ("Total Al2O3_min (Al2O3+1.8895*Al >= 110.00 %) cannot be met by any mix of the allowed raw "
            "materials, even ignoring all other spec limits: the highest achievable Al2O3+1.8895*Al is "
            "107.90 %.") in text
    assert diagnosis.table.set_index("Limit").loc["Total Al2O3_min", "Limits"].iloc[0] == "Al2O3+1.8895*Al"


def test_excel_report_shows_the_expression(write_example, tmp_path):
    from blend_optimizer.pipeline import optimise
    data = load_inputs(write_example(pq_example([("Total Al2O3", TOTAL, None, 80, None, None)])))
    outcome = optimise(data, tmp_path / "report.xlsx", show=False)
    workbook = openpyxl.load_workbook(outcome.excel_path)
    assert workbook["Spec check"].cell(2, 2).value == "Al2O3+1.8895*Al"
    assert workbook["Shadow prices"].cell(2, 3).value == "Al2O3+1.8895*Al"


def test_template_readme_explains_factors(tmp_path):
    from blend_optimizer.templates import write_blank_templates
    single = write_blank_templates(tmp_path)[0]
    readme = " ".join(str(cell.value) for row in openpyxl.load_workbook(single)["README"].iter_rows()
                      for cell in row if cell.value)
    assert "Al2O3+1.8895*Al" in readme and "101.96 / (2 x 26.98)" in readme

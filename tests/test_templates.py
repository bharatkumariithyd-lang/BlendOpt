"""Stage 1 tests: blank templates and the example workbooks."""

import openpyxl
import pytest

import make_templates
from blend_optimizer import config
from blend_optimizer.data import InputError
from blend_optimizer.inputs import load_inputs
from blend_optimizer.templates import write_blank_templates


def _headers(path, sheet):
    worksheet = openpyxl.load_workbook(path)[sheet]
    return [cell.value for cell in worksheet[1]]


def test_blank_templates_have_both_layouts(tmp_path):
    single, master, product = write_blank_templates(tmp_path, components=["X1", "X2"])

    assert openpyxl.load_workbook(single).sheetnames == [
        config.SHEET_README, config.SHEET_RAW_MATERIALS, config.SHEET_SPECS,
        config.SHEET_SETTINGS, config.SHEET_USAGE_LIMITS, config.SHEET_WHAT_IF,
    ]
    assert openpyxl.load_workbook(master).sheetnames == [config.SHEET_README, config.SHEET_RAW_MATERIALS]
    assert openpyxl.load_workbook(product).sheetnames == [
        config.SHEET_README, config.SHEET_SPECS, config.SHEET_SETTINGS, config.SHEET_USAGE_LIMITS,
        config.SHEET_WHAT_IF,
    ]
    assert _headers(single, config.SHEET_WHAT_IF) == list(config.WHAT_IF_COLUMNS)
    assert _headers(single, config.SHEET_RAW_MATERIALS) == list(config.RAW_MATERIAL_FIXED_COLUMNS) + ["X1", "X2"]
    assert _headers(single, config.SHEET_SPECS) == list(config.SPEC_COLUMNS)
    assert _headers(single, config.SHEET_USAGE_LIMITS) == list(config.USAGE_LIMIT_COLUMNS)
    settings = openpyxl.load_workbook(single)[config.SHEET_SETTINGS]
    keys = {row[0].value: row[1].value for row in settings.iter_rows(min_row=2)}
    assert list(keys) == list(config.SETTINGS_KEYS)
    assert keys[config.KEY_SCALE_RESOLUTION] is None  # optional, left blank


def test_blank_template_is_rejected_until_filled_in(tmp_path):
    single = write_blank_templates(tmp_path, components=["X1"])[0]
    with pytest.raises(InputError) as info:
        load_inputs(single)
    assert any("no raw materials listed" in e for e in info.value.errors)


def test_make_templates_writes_loadable_examples(tmp_path):
    written = make_templates.main(tmp_path)
    names = {path.name for path in written}
    assert len(written) == 13
    assert {"Test1_hand_worked.xlsx", "Test2_synthetic_slag_dry.xlsx", "Broken_example.xlsx",
            "Test2_synthetic_slag_with_what_if.xlsx"} <= names

    for path in written:
        if path.name.startswith("Test") and path.name != "Test2_product_only.xlsx":
            assert load_inputs(path).warnings == [], path.name
    split = load_inputs(tmp_path / "Test2_product_only.xlsx", tmp_path / "Example_master_raw_materials.xlsx")
    assert len(split.raw_materials) == 6
    with pytest.raises(InputError):
        load_inputs(tmp_path / "Broken_example.xlsx")

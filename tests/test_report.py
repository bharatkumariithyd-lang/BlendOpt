"""Stage 4 tests: shadow prices, reduced costs, display rounding, Excel and console report."""

import math
from dataclasses import replace

import openpyxl
import pytest

from blend_optimizer import config, examples
from blend_optimizer import solve_and_report as sr
from blend_optimizer.data import UsageLimit
from blend_optimizer.inputs import load_inputs
from blend_optimizer.model import build_model
from blend_optimizer.pipeline import optimise


def _load(write_example, example, file_name="product.xlsx"):
    return load_inputs(write_example(example, file_name))


def _run(data):
    return sr.solve(build_model(data))


def _slag():
    return examples.synthetic_slag_example()


def _with_spec(example, name, **cells):
    """Copy of ``example`` with cells of one Specs row changed, e.g. Min=12."""
    specs = example.specs.copy()
    for column, value in cells.items():
        specs.loc[specs[config.COL_NAME] == name, column] = value
    return replace(example, specs=specs)


def relax(data, result, constraint, delta):
    """Copy of ``data`` with one binding limit relaxed by ``delta`` (Min lowered, Max raised)."""
    if constraint in result.rows:  # a spec row such as "CaO_min"
        row = result.rows[constraint]
        specs = []
        for spec in data.specs:
            if spec.name == row.spec.name and row.bound == "min":
                spec = replace(spec, min_value=spec.min_value - delta)
            elif spec.name == row.spec.name:
                spec = replace(spec, max_value=spec.max_value + delta)
            specs.append(spec)
        return replace(data, specs=specs)

    # A usage limit such as "Raw bauxite max use": set plant and product limit to the relaxed value
    name, bound, _ = constraint.rsplit(" ", 2)
    current = data.effective_limit(name)
    if bound == "max":
        new = UsageLimit(current.min_pct, current.max_pct + delta)
    else:
        new = UsageLimit(current.min_pct - delta, current.max_pct)
    material = replace(data.raw_materials[name], plant_limit=new)
    return replace(data, raw_materials={**data.raw_materials, name: material},
                   product_limits={**data.product_limits, name: new})


# ===========================================================================
# Shadow prices
# ===========================================================================

def test_hand_worked_shadow_prices(write_example):
    data = _load(write_example, examples.hand_worked_example())
    table = sr.shadow_price_table(data, _run(data)).set_index("Constraint")

    assert list(table.index) == ["CaO_min", "Al2O3_min"]  # SiO2_max has slack: not listed
    assert table.loc["CaO_min", "Saving per 1 unit (Rs/dry t)"] == pytest.approx(450.479, abs=0.001)
    assert table.loc["Al2O3_min", "Saving per 1 unit (Rs/dry t)"] == pytest.approx(738.019, abs=0.001)
    assert table.loc["CaO_min", "Relax by"] == "lowering the Min"
    assert table.loc["CaO_min", "Limit unit"] == "wt % (dry basis)"


def test_phosphorus_shadow_price_units(write_example):
    data = _load(write_example, _slag())
    table = sr.shadow_price_table(data, _run(data)).set_index("Constraint")
    p_max = table.loc["P_max"]
    assert p_max["Relax by"] == "raising the Max"
    assert p_max["Step"] == pytest.approx(0.01)
    assert p_max["Saving per step (Rs/dry t)"] == pytest.approx(7350, abs=5)   # per 0.01 % P
    assert p_max["Saving per 1 unit (Rs/dry t)"] == pytest.approx(100 * p_max["Saving per step (Rs/dry t)"])


FINITE_DIFFERENCE_CASES = {
    "test1": lambda: examples.hand_worked_example(),
    "dry": lambda: _slag(),
    "ignited": lambda: examples.synthetic_slag_example(spec_basis="ignited"),
    "as_received": lambda: examples.synthetic_slag_example(spec_basis="as_received"),
    "margin": lambda: examples.synthetic_slag_example(margin=0.5),
    "ratio_min": lambda: _with_spec(_slag(), "Basicity", Min=11.5, Margin=0.5),
    "ratio_max": lambda: _with_spec(_slag(), "Basicity", Max=11.3),
    "usage_max": lambda: replace(_slag(), usage_limits=examples.usage_limit_table([("Raw bauxite", None, 20)])),
    "usage_min": lambda: replace(_slag(), usage_limits=examples.usage_limit_table([("Lime fines", 5, None)])),
}


@pytest.mark.parametrize("case", list(FINITE_DIFFERENCE_CASES))
def test_shadow_prices_match_a_resolve(write_example, case):
    """Move each binding limit a little both ways, re-solve, and compare with the table.

    This checks the sign rules and the pi * sum(g_i x_i) conversion for every
    kind of row: dry, ignited and as-received composition rows, ratio rows
    (Min and Max) and usage limits (Max and Min).

    A central difference (relax by delta AND tighten by delta) is used because
    on ratio and non-dry rows the cost curves slightly as the limit moves (the
    limit sits inside the coefficients); a one-sided difference would pick up
    that curvature, the central one cancels it.
    """
    data = _load(write_example, FINITE_DIFFERENCE_CASES[case]())
    base = _run(data)
    table = sr.shadow_price_table(data, base)
    assert not table.empty
    for _, line in table.iterrows():
        delta = 1e-3 * line["Step"]
        relaxed = _run(relax(data, base, line["Constraint"], delta))
        tightened = _run(relax(data, base, line["Constraint"], -delta))
        saving_by_resolve = (tightened.cost_per_dry_t - relaxed.cost_per_dry_t) / (2 * delta)
        reported = line["Saving per 1 unit (Rs/dry t)"]
        assert reported == pytest.approx(saving_by_resolve, rel=1e-5, abs=0.01), line["Constraint"]
        assert reported >= -1e-6, "relaxing a limit can never cost money"


def test_usage_limit_shadow_rows(write_example):
    example = replace(_slag(), usage_limits=examples.usage_limit_table([("Raw bauxite", None, 20)]))
    data = _load(write_example, example)
    table = sr.shadow_price_table(data, _run(data)).set_index("Constraint")
    row = table.loc["Raw bauxite max use"]
    assert (row["Type"], row["Internal target"], row["Limit unit"]) == ("Usage limit", 20, "wt % of dry blend")
    assert row["Saving per 1 unit (Rs/dry t)"] > 0


def test_relaxation_step():
    assert [sr.relaxation_step(v) for v in (50, 6, 1, 0.5, 0.08, 0.03, 0.005, 0)] == \
           [1, 1, 1, 0.1, 0.01, 0.01, 0.001, 1]


# ===========================================================================
# Reduced costs
# ===========================================================================

def test_reduced_cost_is_the_entry_price(write_example):
    data = _load(write_example, _slag())
    base = _run(data)
    table = sr.reduced_cost_table(data, base).set_index("Raw material")
    assert list(table.index) == ["Lime fines"]

    fines = data.raw_materials["Lime fines"]
    dj = base.reduced_costs["Lime fines"]
    assert table.loc["Lime fines", "Reduced cost (Rs/dry t)"] == pytest.approx(dj)
    # As-received price: dry drop dj is a drop of dj * (1 - 1.0/100) in the quoted price
    assert table.loc["Lime fines", "Price drop needed (Rs/t as entered)"] == pytest.approx(dj * 0.99)
    entry_price = table.loc["Lime fines", "Enters below (Rs/t as entered)"]
    assert entry_price == pytest.approx(fines.cost_input - dj * 0.99)

    for price, should_be_used in ((entry_price + 1, False), (entry_price - 1, True)):
        cheaper = replace(fines, cost_input=price)
        result = _run(replace(data, raw_materials={**data.raw_materials, "Lime fines": cheaper}))
        assert (result.fractions["Lime fines"] > 1e-6) == should_be_used, price


def test_banned_materials(write_example):
    bans = examples.usage_limit_table([("Lime fines", None, 0), ("Calcined dolomite", None, 0)])
    data = _load(write_example, replace(_slag(), usage_limits=bans))
    result = _run(data)
    assert result.is_optimal

    reduced = sr.reduced_cost_table(data, result).set_index("Raw material")
    assert reduced.loc["Lime fines", "Note"] == "Banned by its usage limit (Max 0 %)"
    # Dolomite was in the unrestricted recipe, so its ban costs money
    assert reduced.loc["Calcined dolomite", "Note"].endswith("it would be used if allowed")

    shadow = sr.shadow_price_table(data, result).set_index("Constraint")
    assert shadow.loc["Calcined dolomite max use", "Saving per 1 unit (Rs/dry t)"] > 0
    assert "Lime fines max use" not in shadow.index  # its ban costs nothing


def test_no_unused_materials_gives_empty_table(write_example):
    data = _load(write_example, examples.hand_worked_example())
    result = _run(data)
    assert sr.reduced_cost_table(data, result).empty
    assert sr.report_notes(data, result)[sr.SHEET_REDUCED_COSTS][0] == "None: every raw material is used."


# ===========================================================================
# Display rounding
# ===========================================================================

def test_format_number():
    assert sr.format_number(52.090557, sr.PCT) == "52.09"
    assert sr.format_number(13003.968815, sr.RS) == "13,004"
    assert sr.format_number(510.4567, sr.KG) == "510.5"
    assert sr.format_number(12.26, sr.KG, scale_resolution_kg=0.5) == "12.5"
    assert sr.format_number(12.24, sr.KG, scale_resolution_kg=0.5) == "12.0"
    assert sr.format_number(1234.0, sr.KG, scale_resolution_kg=5) == "1,235.0"
    assert sr.format_number(-1e-12, sr.PCT) == "0.00"     # no "-0.00"
    assert sr.format_number(0.01, sr.PLAIN) == "0.01"
    assert sr.format_number(float("nan"), sr.PCT) == ""
    assert sr.format_number(None, sr.RS) == ""
    assert sr.format_number(math.inf, sr.PCT) == "inf"
    assert sr.format_number("Binding (Min)", sr.PCT) == "Binding (Min)"


def test_column_kinds():
    assert sr.column_kind("Dry blend (wt %)") == sr.PCT
    assert sr.column_kind("As-received mass (kg/batch)") == sr.KG
    assert sr.column_kind("Saving per step (Rs/dry t)") == sr.RS
    assert sr.column_kind("Step") == sr.PLAIN


def test_tables_are_not_rounded(write_example):
    data = _load(write_example, _slag())
    recipe = sr.recipe_table(data, _run(data)).set_index("Raw material")
    lime = recipe.loc["Lime", "Dry blend (wt %)"]
    assert lime != round(lime, 2)  # full precision until displayed


def test_display_table_has_blanks_not_nan(write_example):
    data = _load(write_example, _slag())
    shown = sr.display_table(sr.spec_check_table(data, _run(data))).set_index("Spec")
    assert shown.loc["SiO2", "Customer Min"] == ""
    assert shown.loc["SiO2", "Customer Max"] == "6.00"
    assert not shown.isin(["nan", "NaN", "None"]).any().any()


# ===========================================================================
# Excel and console report
# ===========================================================================

def test_excel_report(write_example, tmp_path):
    data = _load(write_example, _slag())
    outcome = optimise(data, tmp_path / "out" / "report.xlsx", show=False)
    assert outcome.excel_path.is_file()
    assert outcome.lp_path == tmp_path / "out" / "report.lp" and outcome.lp_path.is_file()

    workbook = openpyxl.load_workbook(outcome.excel_path)
    assert workbook.sheetnames == ["Summary", "Recipe", "Composition", "Spec check", "Cost",
                                   "Shadow prices", "Reduced costs"]
    recipe = workbook["Recipe"]
    headers = [cell.value for cell in recipe[1]]
    wt, dry_kg = headers.index("Dry blend (wt %)") + 1, headers.index("Dry mass (kg/batch)") + 1
    lime_row, total_row = 2, 8
    assert recipe.cell(total_row, 1).value == "TOTAL"
    # wt %: full precision stored, shown with 2 decimals
    assert recipe.cell(lime_row, wt).value == pytest.approx(100 * outcome.result.fractions["Lime"])
    assert recipe.cell(lime_row, wt).number_format == "0.00"
    # kg: stored rounded to the 0.1 kg scale, and the rows add up exactly to the TOTAL row
    dry_rows = [recipe.cell(row, dry_kg).value for row in range(2, total_row)]
    assert all(round(value * 10) == pytest.approx(value * 10) for value in dry_rows)
    assert round(sum(dry_rows), 6) == recipe.cell(total_row, dry_kg).value == 980.1
    assert recipe.cell(lime_row, dry_kg).number_format == "#,##0.0"

    spec = workbook["Spec check"]
    spec_headers = [cell.value for cell in spec[1]]
    sio2_row = 1 + [cell.value for cell in spec["A"]][1:].index("SiO2") + 1
    assert spec.cell(sio2_row, spec_headers.index("Customer Min") + 1).value is None  # blank, not NaN

    cost = workbook["Cost"]
    assert cost.cell(2, 2).number_format == "#,##0"

    composition = workbook["Composition"]
    formats = {row[0].value: row[2].number_format for row in composition.iter_rows(min_row=2) if row[2].value is not None}
    assert (formats["P"], formats["CaO"], formats["TiO2"]) == ("0.000", "0.00", "0.00")

    notes = [cell.value for cell in workbook["Shadow prices"]["A"] if cell.value]
    assert any(text.startswith("Valid only near the current limit") for text in notes)


def test_excel_scale_resolution(write_example, tmp_path):
    data = _load(write_example, _slag())
    result = _run(data)
    path = sr.write_excel(sr.build_report(data, result), tmp_path / "r.xlsx", scale_resolution_kg=0.5)
    recipe = openpyxl.load_workbook(path)["Recipe"]
    for row in range(2, 9):
        value = recipe.cell(row, 3).value  # Dry mass (kg/batch)
        assert value * 2 == round(value * 2), value


def test_console_summary(write_example):
    data = _load(write_example, examples.hand_worked_example())
    text = sr.format_summary(data, _run(data))
    assert "Cost: Rs 11,482 per dry t of blend" in text
    assert "Lime            53.40               534.0" in text
    assert "CaO_min lowering the Min" in text and "450" in text and "738" in text
    assert "valid only near the current limit" in text
    assert "nan" not in text.lower()


def test_infeasible_report(write_example, tmp_path, capsys):
    data = _load(write_example, examples.hand_worked_example(al2o3_min=45))
    outcome = optimise(data, tmp_path / "inf.xlsx")
    assert list(outcome.tables) == ["Summary", "Diagnosis"]  # Stage 6 adds the diagnosis
    assert openpyxl.load_workbook(outcome.excel_path).sheetnames == ["Summary", "Diagnosis"]
    assert outcome.lp_path.is_file()
    assert "No recipe: the solver status is 'Infeasible'." in capsys.readouterr().out


# ===========================================================================
# Stage 4 follow-up: data-driven spec decimals, weigh sheet, scale resolution
# ===========================================================================

def test_decimals_entered_and_limit_decimals():
    assert [sr.decimals_entered(v) for v in (50, 50.0, 0.03, 0.125, 1e-05)] == [0, 0, 2, 3, 5]
    assert sr.decimals_entered(0.1 + 0.2) == config.DISPLAY_DECIMALS_PCT_MAX  # computed noise is capped
    assert sr.limit_decimals([50, 56, 0]) == 2          # minimum 2
    assert sr.limit_decimals([None, 0.03, 0]) == 3      # one more than 0.03
    assert sr.limit_decimals([float("nan"), 0.015]) == 4
    assert sr.limit_decimals([]) == 2


def test_spec_rows_show_one_more_decimal_than_entered(write_example):
    data = _load(write_example, _slag())
    result = _run(data)
    check = sr.display_table(sr.spec_check_table(data, result)).set_index("Spec")
    assert (check.loc["P", "Customer Max"], check.loc["P", "Target Max"], check.loc["P", "Achieved"]) ==            ("0.030", "0.030", "0.030")
    assert (check.loc["S", "Target Max"], check.loc["S", "Achieved"]) == ("0.080", "0.040")
    assert (check.loc["CaO", "Target Min"], check.loc["CaO", "Achieved"]) == ("50.00", "50.00")
    assert check.loc["Basicity", "Achieved"] == "11.40"
    # Spec-limited components follow the same rule in the composition table; others keep 2 dp
    composition = sr.display_table(sr.composition_table(data, result),
                                   row_decimals=sr.composition_decimals(data)).set_index("Component")
    assert composition.loc["P", "Blend (wt %)"] == "0.030"
    assert composition.loc["S", "Blend (wt %)"] == "0.040"
    assert composition.loc["CaO", "Blend (wt %)"] == "50.00"
    for unlimited in ("TiO2", "LOI", "Sum of components"):  # no spec of their own: default 2 dp
        assert len(composition.loc[unlimited, "Blend (wt %)"].split(".")[1]) == 2, unlimited
    # The shadow price table shows the same limit with the same decimals
    shadow = sr.display_table(sr.shadow_price_table(data, result)).set_index("Constraint")
    assert shadow.loc["P_max", "Customer limit"] == "0.030"


def test_margin_decimals_count(write_example):
    example = _with_spec(examples.hand_worked_example(), "CaO", Margin=0.125)
    data = _load(write_example, example)
    check = sr.display_table(sr.spec_check_table(data, _run(data))).set_index("Spec")
    assert check.loc["CaO", "Target Min"] == "50.1250"   # 0.125 has 3 decimals -> 4 shown
    assert check.loc["SiO2", "Target Max"] == "6.00"     # other rows keep their own decimals


def test_weigh_sheet_rounding():
    rounded, total, index, difference = sr.weigh_sheet_rounding([0.34, 0.33, 0.33], 0.1)
    assert (rounded, total, index, difference) == ([0.4, 0.3, 0.3], 1.0, 0, 0.1)
    rounded, total, index, difference = sr.weigh_sheet_rounding([510.46, 469.64], 0.1)
    assert (rounded, total, index, difference) == ([510.5, 469.6], 980.1, None, 0.0)


def test_weigh_sheet_puts_difference_on_largest(write_example):
    data = _load(write_example, _slag())
    recipe = sr.recipe_table(data, _run(data))
    sheet, adjustments = sr.weigh_sheet(recipe, 5.0)
    rows = sheet.set_index("Raw material")["As-received mass (kg/batch)"]
    # Rounded singly to 5 kg the rows add up to 1005; the -5 kg goes on Lime (the largest)
    assert rows["Lime"] == 510.0
    assert rows.drop("TOTAL").sum() == rows["TOTAL"] == 1000.0
    by_column = {adjustment.column: adjustment for adjustment in adjustments}
    assert by_column["As-received mass (kg/batch)"].material == "Lime"
    assert by_column["As-received mass (kg/batch)"].difference_kg == -5.0
    notes = sr.weigh_sheet_notes(adjustments, 5.0)
    assert "As-received mass (kg/batch): rounding difference of -5.0 kg put on Lime (the largest)." in notes
    # The table itself is not changed: rounding is for display only
    assert recipe.set_index("Raw material").loc["Lime", "As-received mass (kg/batch)"] != 510.0


def test_scale_resolution_from_settings(write_example, tmp_path):
    example = examples.synthetic_slag_example()
    example = replace(example, settings={**example.settings, config.KEY_SCALE_RESOLUTION: 5})
    data = _load(write_example, example)
    assert data.settings.scale_resolution_kg == 5

    outcome = optimise(data, tmp_path / "scale5.xlsx", show=False)
    recipe = openpyxl.load_workbook(outcome.excel_path)["Recipe"]
    headers = [cell.value for cell in recipe[1]]
    ar = headers.index("As-received mass (kg/batch)") + 1
    rows = [recipe.cell(row, ar).value for row in range(2, 8)]
    assert all(value % 5 == 0 for value in rows)
    assert sum(rows) == recipe.cell(8, ar).value == 1000
    notes = [cell.value for cell in recipe["A"] if cell.value]
    assert any("-5.0 kg put on Lime (the largest)" in str(note) for note in notes)
    assert "Scale resolution" in outcome.tables["Summary"]["Item"].tolist()


def test_console_shows_weigh_sheet_note(write_example):
    data = _load(write_example, _slag())
    text = sr.format_summary(data, _run(data))
    assert "each kg column adds up exactly to its TOTAL" in text
    assert "0.030" in text  # P target and achieved with 3 decimals

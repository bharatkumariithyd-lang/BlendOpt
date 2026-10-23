"""Stage 7 tests: the WhatIf sheet, run.py and the full report workbook."""

from dataclasses import replace

import openpyxl
import pytest

import run as run_module
from blend_optimizer import config, examples
from blend_optimizer.data import InputError
from blend_optimizer.inputs import load_inputs
from blend_optimizer.pipeline import diagnose_infeasible
from blend_optimizer.what_if import check_scenarios, scenario_sheet_name, what_if


def _with_what_if(example, rows):
    return replace(example, what_if=examples.what_if_table(rows))


def _slag_with(rows):
    return _with_what_if(examples.synthetic_slag_example(), rows)


# ===========================================================================
# Reading the WhatIf sheet
# ===========================================================================

def test_what_if_sheet_becomes_scenarios(write_example):
    data = load_inputs(write_example(examples.synthetic_slag_example(with_what_if=True)))
    by_name = {scenario.name: scenario for scenario in data.scenarios}
    assert list(by_name) == ["Fluorspar dearer", "Looser P", "Capped bauxite", "No fluorspar"]
    assert by_name["Fluorspar dearer"].price_changes == {"Fluorspar": 35000}
    assert by_name["Looser P"].spec_changes == {"P": {"Max": 0.035}}
    # Two rows with the same Scenario name make one scenario
    assert by_name["Capped bauxite"].usage_changes == {"Raw bauxite": {"Max": 20}}
    assert by_name["Capped bauxite"].spec_changes == {"CaO": {"Margin": 0.5}}
    assert data.sources[config.SHEET_WHAT_IF].endswith("product.xlsx")
    assert check_scenarios(data) == []


def test_what_if_blank_values_and_lenient_types(write_example):
    data = load_inputs(write_example(_slag_with([
        ("Open CaO", "spec max", "CaO", None),        # blank Max removes that limit
        ("No margin", "Spec Margin", "CaO", None),    # blank Margin means 0
        ("Reset", "USAGE MAX", "Lime fines", None),   # blank resets to the plant limit
    ])))
    by_name = {scenario.name: scenario for scenario in data.scenarios}
    assert by_name["Open CaO"].spec_changes == {"CaO": {"Max": None}}
    assert by_name["No margin"].spec_changes == {"CaO": {"Margin": None}}
    assert by_name["Reset"].usage_changes == {"Lime fines": {"Max": None}}
    assert check_scenarios(data) == []


WHAT_IF_ERRORS = [
    pytest.param([("S", "Spec Minimum", "CaO", 51)], ["'Spec Minimum' is not known", "Did you mean 'Spec Min'?"],
                 id="unknown-change-type"),
    pytest.param([("S", "Price", "Flourspar", 30000)], ["'Flourspar' is not a raw material", "Did you mean 'Fluorspar'?"],
                 id="unknown-material"),
    pytest.param([("S", "Spec Max", "Phosphorus", 0.04)], ["'Phosphorus' is not a spec in Specs"], id="unknown-spec"),
    pytest.param([(None, "Price", "Lime", 9500)], ["Scenario is blank"], id="blank-scenario"),
    pytest.param([("S", "Price", None, 9500)], ["Item is blank"], id="blank-item"),
    pytest.param([("S", "Price", "Lime", None)], ["New value: value is blank"], id="price-blank"),
    pytest.param([("S", "Price", "Lime", "dear")], ["'dear' is not a number"], id="price-text"),
    pytest.param([("S", "Usage Max", "Lime", 120)], ["120 is above the maximum allowed (100)"], id="usage-above-100"),
    pytest.param([("S", "Spec Min", "CaO", -1)], ["-1 is negative"], id="negative"),
    pytest.param([("S", "Price", "Lime", 9500), ("S", "Price", "Lime", 9600)], ["already given in row 2"],
                 id="duplicate-row"),
]


@pytest.mark.parametrize("rows, fragments", WHAT_IF_ERRORS)
def test_what_if_row_errors(write_example, rows, fragments):
    with pytest.raises(InputError) as info:
        load_inputs(write_example(_slag_with(rows)))
    assert len(info.value.errors) == 1, info.value.errors
    for fragment in fragments:
        assert fragment in info.value.errors[0]


def test_what_if_missing_column(write_example, tmp_path):
    example = _slag_with([("S", "Price", "Lime", 9500)])
    example = replace(example, what_if=example.what_if.drop(columns=[config.COL_ITEM]))
    with pytest.raises(InputError) as info:
        load_inputs(write_example(example))
    assert "required column 'Item' is missing" in info.value.errors[0]


def test_scenario_consistency_is_checked_before_solving(write_example):
    data = load_inputs(write_example(_slag_with([
        ("Too high", "Spec Min", "CaO", 57),          # above CaO Max 56
        ("Nothing left", "Spec Max", "SiO2", None),   # SiO2 has only a Max
        ("Fine", "Price", "Lime", 9500),
    ])))
    errors = check_scenarios(data)
    assert errors == [
        "WhatIf scenario 'Too high': Spec 'CaO': Min (57) is greater than Max (56).",
        "WhatIf scenario 'Nothing left': Spec 'SiO2': give a Min, a Max, or both.",
    ]


# ===========================================================================
# run(): base case + scenarios in one workbook
# ===========================================================================

def test_full_run_with_scenarios(write_example, tmp_path):
    product = write_example(examples.synthetic_slag_example(with_what_if=True))
    full_run = run_module.run(product, output_path=tmp_path / "report.xlsx", show=False)

    workbook = openpyxl.load_workbook(full_run.base.excel_path)
    assert workbook.sheetnames == [
        "Summary", "Recipe", "Composition", "Spec check", "Cost", "Shadow prices", "Reduced costs",
        "Price ranges", "Scenarios", "Scenario Fluorspar dearer", "Scenario Looser P",
        "Scenario Capped bauxite", "Scenario No fluorspar",
    ]
    assert full_run.base.lp_path.is_file()

    # Overview: one line per scenario
    overview = workbook["Scenarios"]
    rows = {row[0]: row for row in overview.iter_rows(min_row=2, values_only=True)}
    headers = [cell.value for cell in overview[1]]
    cost, status = headers.index("Cost (Rs/dry t)"), headers.index("Status")
    assert rows["No fluorspar"][status] == "Infeasible" and rows["No fluorspar"][cost] is None
    assert rows["Looser P"][headers.index("Changes")] == "Spec Max P 0.0300 -> 0.0350"

    # Each scenario's cost equals a direct what-if with the same changes
    data = full_run.base.data
    direct = what_if(data, price_changes={"Fluorspar": 35000})["Cost"].set_index("Item")
    assert rows["Fluorspar dearer"][cost] == pytest.approx(
        direct.loc["Cost (Rs per dry tonne of blend)", "Scenario"])

    # A scenario sheet stacks its tables under bold titles
    titles = [cell.value for cell in workbook["Scenario Looser P"]["A"]
              if cell.value and str(cell.value).startswith("Looser P:")]
    assert titles == ["Looser P: changes", "Looser P: cost", "Looser P: recipe", "Looser P: spec check",
                      "Looser P: binding limits"]


def test_infeasible_scenario_gets_the_same_diagnosis_as_a_base_case(write_example, tmp_path):
    product = write_example(examples.synthetic_slag_example(with_what_if=True))
    full_run = run_module.run(product, output_path=tmp_path / "report.xlsx", show=False)
    no_fluorspar = next(outcome for outcome in full_run.scenarios if outcome.scenario.name == "No fluorspar")
    assert no_fluorspar.result.status == "Infeasible"

    # The scenario is named as the source of the ban ...
    label = "WhatIf (No fluorspar)"
    assert any(f"Fluorspar (CaF2 85.00 %) is not allowed (Max 0 %, {label})" in message
               for message in no_fluorspar.diagnosis.messages)
    # ... and apart from that source label, the messages are those of the same case as a base case
    # (Fluorspar banned through the UsageLimits sheet)
    banned = load_inputs(write_example(examples.synthetic_slag_example(ban_fluorspar=True), "ban.xlsx"))
    base_messages = diagnose_infeasible(banned).messages
    assert any("UsageLimits" in message for message in base_messages)
    without_label = [message.replace(label, "UsageLimits") for message in no_fluorspar.diagnosis.messages]
    assert without_label == base_messages

    sheet = openpyxl.load_workbook(full_run.base.excel_path)["Scenario No fluorspar"]
    column_a = [str(cell.value) for cell in sheet["A"] if cell.value]
    assert "No fluorspar: diagnosis" in column_a
    assert any("No allowed raw material contains CaF2, so CaF2_min" in text for text in column_a)
    # The feasible scenarios have no diagnosis section
    assert "Looser P: diagnosis" not in [cell.value for cell in openpyxl.load_workbook(
        full_run.base.excel_path)["Scenario Looser P"]["A"]]


def test_scenario_that_fixes_an_infeasible_base(write_example, tmp_path):
    example = _with_what_if(examples.hand_worked_example(al2o3_min=45), [("Al2O3 39", "Spec Min", "Al2O3", 39)])
    full_run = run_module.run(write_example(example), output_path=tmp_path / "fix.xlsx", show=False)
    assert full_run.base.result.status == "Infeasible" and full_run.base.diagnosis is not None
    assert full_run.scenarios[0].result.is_optimal
    sheets = openpyxl.load_workbook(full_run.base.excel_path).sheetnames
    assert sheets == ["Summary", "Diagnosis", "Price ranges", "Scenarios", "Scenario Al2O3 39"]


def test_invalid_scenario_stops_the_run(write_example, tmp_path):
    product = write_example(_slag_with([("Too high", "Spec Min", "CaO", 57)]))
    with pytest.raises(InputError) as info:
        run_module.run(product, output_path=tmp_path / "never.xlsx", show=False)
    assert "WhatIf scenario 'Too high'" in info.value.errors[0]
    assert not (tmp_path / "never.xlsx").exists()


def test_run_without_what_if_sheet(write_example, tmp_path):
    full_run = run_module.run(write_example(examples.hand_worked_example()), output_path=tmp_path / "t1.xlsx",
                              show=False)
    assert full_run.scenarios == []
    sheets = openpyxl.load_workbook(full_run.base.excel_path).sheetnames
    assert "Scenarios" not in sheets
    assert sheets[-1] == "Price ranges"  # every report has the base-case price ranges


# ===========================================================================
# The Price ranges sheet
# ===========================================================================

def test_price_ranges_sheet(write_example, tmp_path, capsys):
    full_run = run_module.run(write_example(examples.hand_worked_example()), output_path=tmp_path / "t1.xlsx")
    printed = capsys.readouterr().out
    assert "Price ranges (base case)" in printed

    table = full_run.price_ranges.set_index("Raw material")
    assert table.loc["Calcined alumina", "Unchanged from (Rs/t)"] == pytest.approx(13271.2, abs=1.0)
    assert table.loc["Bauxite", "Unchanged up to (Rs/t)"] == pytest.approx(23750.0, abs=1.0)

    sheet = openpyxl.load_workbook(full_run.base.excel_path)["Price ranges"]
    headers = [cell.value for cell in sheet[1]]
    rows = {row[0]: row for row in sheet.iter_rows(min_row=2, max_row=4, values_only=True)}
    assert rows["Calcined alumina"][headers.index("Unchanged from (Rs/t)")] == pytest.approx(13271.2, abs=1.0)
    assert sheet.cell(2, headers.index("Current price (Rs/t)") + 1).number_format == "#,##0"
    notes = [str(cell.value) for cell in sheet["A"] if cell.value]
    assert any("accurate to Rs 1/t" in note for note in notes)
    assert any("0.01 wt %" in note for note in notes)


def test_price_ranges_sheet_for_an_infeasible_base(write_example, tmp_path):
    full_run = run_module.run(write_example(examples.hand_worked_example(al2o3_min=45)),
                              output_path=tmp_path / "inf.xlsx", show=False)
    assert full_run.price_ranges.empty
    sheet = openpyxl.load_workbook(full_run.base.excel_path)["Price ranges"]
    notes = [str(cell.value) for cell in sheet["A"] if cell.value]
    assert any(note.startswith("Not available: price ranges need a feasible base case") for note in notes)


def test_scenario_sheet_names():
    used = {"summary", "recipe"}
    assert scenario_sheet_name("Recipe", used) == "Scenario Recipe"
    assert scenario_sheet_name("P/S: tighter?", used) == "Scenario P S  tighter"
    long_name = "A very long scenario name that Excel cannot take"
    first = scenario_sheet_name(long_name, used)
    second = scenario_sheet_name(long_name, used)
    assert len(first) == 31 and len(second) == 31 and first != second and second.endswith("(2)")


# ===========================================================================
# The command line
# ===========================================================================

def test_main_writes_the_report(write_example, tmp_path, capsys):
    product = write_example(examples.hand_worked_example(), "test1.xlsx")
    exit_code = run_module.main([str(product), "--output", str(tmp_path / "out.xlsx")])
    printed = capsys.readouterr().out
    assert exit_code == 0
    assert "Cost: Rs 11,482 per dry t of blend" in printed
    assert f"Report written to: {tmp_path / 'out.xlsx'}" in printed
    assert (tmp_path / "out.xlsx").is_file() and (tmp_path / "out.lp").is_file()


def test_main_with_master_file_and_quiet(write_example, tmp_path, capsys):
    master, product = examples.synthetic_slag_example().split()
    exit_code = run_module.main([str(write_example(product, "product.xlsx")),
                                 "--raw-materials", str(write_example(master, "master.xlsx")),
                                 "--output", str(tmp_path / "split.xlsx"), "--quiet"])
    printed = capsys.readouterr().out
    assert exit_code == 0
    assert "Least-cost blend" not in printed      # --quiet: no summary
    assert "Report written to" in printed


def test_main_reports_input_errors(write_example, tmp_path, capsys):
    exit_code = run_module.main([str(write_example(examples.broken_example(), "broken.xlsx")),
                                 "--output", str(tmp_path / "broken_result.xlsx")])
    printed = capsys.readouterr().out
    assert exit_code == 1
    assert "9 input error(s) found:" in printed
    assert "1 warning(s):" in printed
    assert "Nothing was solved" in printed
    assert not (tmp_path / "broken_result.xlsx").exists()


def test_main_explains_a_locked_report(write_example, tmp_path, capsys, monkeypatch):
    def locked(*args, **kwargs):
        raise PermissionError("[Errno 13] Permission denied: 'out.xlsx'")

    monkeypatch.setattr(run_module, "optimise", locked)
    exit_code = run_module.main([str(write_example(examples.hand_worked_example())), "--quiet"])
    assert exit_code == 1
    assert "If the report is open in Excel, close it and run again." in capsys.readouterr().out


def test_default_output_path():
    assert run_module.default_output_path("data/My product.xlsx") == \
           run_module.DEFAULT_OUTPUT_FOLDER / "My product_result.xlsx"

"""Solve the blend LP and turn the answer into results, report tables,
an Excel workbook and a console summary.

(PuLP lives here and in ``model``; every other module works with the plain
``SolveResult`` returned by :func:`solve`.)

Sections:
    1. Solver choice
    2. Solving -> SolveResult (plain Python values)
    3. Blend masses
    4. Report tables: recipe, composition, spec check, cost, summary
    5. Sensitivity: shadow prices and reduced costs
    6. Display: rounding, Excel workbook, console summary

The tables in sections 4 and 5 keep full precision. Numbers are rounded only
when shown (section 6), as set by ``config.DISPLAY_*`` and
``config.SCALE_RESOLUTION_KG``.
"""

import math
from dataclasses import dataclass, field
from decimal import Decimal
from pathlib import Path
from typing import Iterable, Sequence

import pandas as pd
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet
from pulp import HiGHS, LpSolver, LpSolveStats, LpSolveStatus

from blend_optimizer import config
from blend_optimizer.data import BlendData, Term, price_from_dry_basis
from blend_optimizer.linear_rows import (
    MIN,
    LinearRow,
    achieved_value,
    basis_factor,
    blend_value,
    row_activity,
)
from blend_optimizer.model import SUM_ROW_NAME, BlendModel
from blend_optimizer.templates import style_header_row

# Values closer to zero than this are treated as exactly zero (solver round-off)
_ZERO = 1e-9

# Report sheet names, in the order they appear in the workbook
SHEET_SUMMARY = "Summary"
SHEET_RECIPE = "Recipe"
SHEET_COMPOSITION = "Composition"
SHEET_SPEC_CHECK = "Spec check"
SHEET_COST = "Cost"
SHEET_SHADOW_PRICES = "Shadow prices"
SHEET_REDUCED_COSTS = "Reduced costs"


# ===========================================================================
# 1. Solver choice: the ONLY place a solver is named
# ===========================================================================

def make_solver(msg: bool = False) -> LpSolver:
    """Return the solver used for every solve in this tool.

    To switch solver later, change only this function.
    (PuLP 4 has no usable default solver here: ``LpSolverDefault`` is None
    unless CBC or GLPK is installed, so the solver is always passed in.)

    Args:
        msg: Show the solver's log on the console.
    """
    return HiGHS(msg=msg)


# ===========================================================================
# 2. Solving
# ===========================================================================

@dataclass
class SolveResult:
    """The answer from one solve, as plain Python values (no PuLP objects).

    Row dicts are keyed by our row names (e.g. "CaO_min", and
    ``model.SUM_ROW_NAME``); material dicts by raw material name.

    Attributes:
        status: Solver status, e.g. "Optimal", "Infeasible", "Unbounded".
        is_optimal: True only if an optimal solution was found.
        cost_per_dry_t: Objective, Rs per dry tonne of blend (None if not optimal).
        fractions: Material -> x_i, dry mass fraction (0..1).
        row_activity: Row -> sum(a_i x_i), the left-hand side.
        row_slack: Row -> slack, >= 0 when satisfied (activity - rhs for >=,
            rhs - activity for <=), in the row's own unit.
        row_dual: Row -> PuLP's ``pi`` = d(cost)/d(rhs). Turned into
            "Rs/t saved per unit of relaxation" by :func:`shadow_price_table`.
        reduced_costs: Material -> PuLP's ``dj``, Rs/dry t per unit fraction.
        rows: Row name -> the LinearRow it was built from (spec rows only;
            plain Python, so the report can recompute achieved values).
        lp_file: Path of the .lp file written, if any.
        elastic: True if this is the answer of the elastic (diagnosis) model.
        violations: Elastic model only: row name -> violation v_k, in the
            row's LP units (0 = the limit is met).
        relative_violation: Elastic model only: the objective,
            sum(v_k / scale_k). (``cost_per_dry_t`` is then the cost of the
            elastic recipe, for information.)
    """

    status: str
    is_optimal: bool
    cost_per_dry_t: float | None = None
    fractions: dict[str, float] = field(default_factory=dict)
    row_activity: dict[str, float] = field(default_factory=dict)
    row_slack: dict[str, float] = field(default_factory=dict)
    row_dual: dict[str, float] = field(default_factory=dict)
    reduced_costs: dict[str, float] = field(default_factory=dict)
    rows: dict[str, LinearRow] = field(default_factory=dict)
    lp_file: str | None = None
    elastic: bool = False
    violations: dict[str, float] = field(default_factory=dict)
    relative_violation: float | None = None


def solve(model: BlendModel, lp_path: str | Path | None = None, msg: bool = False) -> SolveResult:
    """Solve the model and copy the answer into a :class:`SolveResult`.

    Everything is copied out of PuLP straight away, because PuLP 4 variables
    and constraints stop working once their problem is gone.

    Args:
        model: From ``model.build_model``.
        lp_path: If given, the human-readable .lp file is written here
            (before solving, so it exists even if the problem is infeasible).
        msg: Show the solver log.

    Returns:
        The result. If the problem is not solved to optimality, only
        ``status``, ``rows`` and ``lp_file`` are filled in.
    """
    lp_file = None
    if lp_path is not None:
        lp_path = Path(lp_path)
        lp_path.parent.mkdir(parents=True, exist_ok=True)
        model.problem.writeLP(str(lp_path))
        lp_file = str(lp_path)

    stats = model.problem.solve(make_solver(msg=msg))
    if stats.status != LpSolveStatus.Optimal or not stats.has_solution:
        return SolveResult(status=stats.status_str, is_optimal=False, rows=dict(model.rows), lp_file=lp_file,
                           elastic=model.elastic)
    return _read_solution(model, stats, lp_file)


def _clean(value: float | None) -> float:
    """Solver round-off such as -3e-17 becomes 0.0."""
    if value is None:
        return 0.0
    return 0.0 if abs(value) < _ZERO else float(value)


def _read_solution(model: BlendModel, stats: LpSolveStats, lp_file: str | None) -> SolveResult:
    """Copy values, slacks, duals and reduced costs out of PuLP."""
    fractions = {name: _clean(variable.varValue) for name, variable in model.fractions.items()}
    reduced_costs = {name: _clean(variable.dj) for name, variable in model.fractions.items()}

    row_activity_: dict[str, float] = {SUM_ROW_NAME: sum(fractions.values())}
    for row_name, row in model.rows.items():
        row_activity_[row_name] = row_activity(row, fractions)

    # In the elastic model the objective is the violation, not the cost
    if model.elastic:
        cost = sum(model.data.raw_materials[name].dry_cost * x for name, x in fractions.items())
        relative_violation = float(stats.objective)
    else:
        cost, relative_violation = float(stats.objective), None

    return SolveResult(
        status=stats.status_str,
        is_optimal=True,
        cost_per_dry_t=cost,
        fractions=fractions,
        row_activity=row_activity_,
        row_slack={name: _clean(c.slack) for name, c in model.constraints.items()},
        row_dual={name: _clean(c.pi) for name, c in model.constraints.items()},
        reduced_costs=reduced_costs,
        rows=dict(model.rows),
        lp_file=lp_file,
        elastic=model.elastic,
        violations={name: _clean(variable.varValue) for name, variable in model.violations.items()},
        relative_violation=relative_violation,
    )


# ===========================================================================
# 3. Blend masses (plain arithmetic on the recipe)
# ===========================================================================

def as_received_tonnes_per_dry_tonne(data: BlendData, fractions: dict[str, float]) -> float:
    """W = sum(x_i / (1 - moisture_i/100)): as-received tonnes per dry tonne of blend."""
    return sum(x / (1.0 - data.raw_materials[name].moisture_pct / 100.0) for name, x in fractions.items())


def batch_dry_kg(data: BlendData, fractions: dict[str, float]) -> float:
    """Dry mass in one batch. Batch_size_kg is AS-RECEIVED, so dry mass = batch / W."""
    return data.settings.batch_size_kg / as_received_tonnes_per_dry_tonne(data, fractions)


# ===========================================================================
# 4. Report tables (full precision)
# ===========================================================================

def _require_solution(result: SolveResult) -> None:
    if not result.is_optimal:
        raise ValueError(f"No solution to report: the solver status is '{result.status}'.")


def recipe_table(data: BlendData, result: SolveResult) -> pd.DataFrame:
    """The recipe: dry wt %, dry and as-received kg per batch, and cost shares.

    The batch is the total AS-RECEIVED mass weighed into the mixer, so the
    as-received column adds up to Batch_size_kg and the dry column to less
    (by the moisture). The last row is the total.
    """
    _require_solution(result)
    dry_total_kg = batch_dry_kg(data, result.fractions)
    records = []
    for name, x in result.fractions.items():
        material = data.raw_materials[name]
        dry_kg = x * dry_total_kg
        records.append({
            "Raw material": name,
            "Dry blend (wt %)": 100.0 * x,
            "Dry mass (kg/batch)": dry_kg,
            "As-received mass (kg/batch)": dry_kg / (1.0 - material.moisture_pct / 100.0),
            "Moisture (%)": material.moisture_pct,
            "Price as entered (Rs/t)": material.cost_input,
            "Price basis": material.price_basis,
            "Dry cost (Rs/dry t)": material.dry_cost,
            "Cost share (Rs/dry t of blend)": material.dry_cost * x,
        })
    table = pd.DataFrame(records)
    total = {
        "Raw material": "TOTAL",
        "Dry blend (wt %)": table["Dry blend (wt %)"].sum(),
        "Dry mass (kg/batch)": table["Dry mass (kg/batch)"].sum(),
        "As-received mass (kg/batch)": table["As-received mass (kg/batch)"].sum(),
        "Cost share (Rs/dry t of blend)": table["Cost share (Rs/dry t of blend)"].sum(),
    }
    return pd.concat([table, pd.DataFrame([total])], ignore_index=True)


def composition_table(data: BlendData, result: SolveResult) -> pd.DataFrame:
    """Blend composition: every component on the spec basis, plus LOI and moisture.

    A component's blend value is sum(c_i x_i) / sum(g_i x_i), with g_i from
    ``linear_rows.basis_factor`` (g_i = 1 on a dry basis).
    """
    _require_solution(result)
    basis = data.settings.spec_basis
    fractions = result.fractions
    denominator = {name: basis_factor(material, basis) for name, material in data.raw_materials.items()}

    records = []
    for component in data.components:
        numerator = {name: material.composition[component] for name, material in data.raw_materials.items()}
        records.append({"Component": component, "Basis": basis,
                        "Blend (wt %)": blend_value(numerator, denominator, fractions)})
    records.append({"Component": "Sum of components", "Basis": basis,
                    "Blend (wt %)": sum(record["Blend (wt %)"] for record in records)})

    loi = {name: material.loi_pct for name, material in data.raw_materials.items()}
    ones = {name: 1.0 for name in data.raw_materials}
    records.append({"Component": "LOI", "Basis": config.BASIS_DRY,
                    "Blend (wt %)": blend_value(loi, ones, fractions)})
    as_received = as_received_tonnes_per_dry_tonne(data, fractions)
    records.append({"Component": "Moisture", "Basis": config.BASIS_AS_RECEIVED,
                    "Blend (wt %)": 100.0 * (as_received - 1.0) / as_received})
    return pd.DataFrame(records)


def spec_check_table(data: BlendData, result: SolveResult) -> pd.DataFrame:
    """One line per Specs row: limits, targets, achieved value, Binding/Slack.

    * Status comes from the solver's slack on each LP row: a row whose slack
      is below ``config.BINDING_TOLERANCE`` is Binding (the limit is
      holding the recipe where it is).
    * "Slack to target" is in the spec's own unit: how far the achieved
      value is from the nearest internal target.
    """
    _require_solution(result)
    rows_by_spec: dict[str, list[LinearRow]] = {}
    for row in result.rows.values():
        rows_by_spec.setdefault(row.spec.name, []).append(row)

    records = []
    for spec in data.specs:
        spec_rows = rows_by_spec[spec.name]
        # Both bounds of a spec limit the same blend value, so either row will do
        achieved = achieved_value(spec_rows[0], result.fractions)
        binding = [row.bound.capitalize() for row in spec_rows
                   if result.row_slack[row.name] < config.BINDING_TOLERANCE]
        distances = []
        if spec.min_target is not None:
            distances.append(achieved - spec.min_target)
        if spec.max_target is not None:
            distances.append(spec.max_target - achieved)
        records.append({
            "Spec": spec.name,
            "Limits": spec.expression,
            "Basis": "ratio (any basis)" if spec.is_ratio else data.settings.spec_basis,
            "Customer Min": spec.min_value,
            "Customer Max": spec.max_value,
            "Margin": spec.margin,
            "Target Min": spec.min_target,
            "Target Max": spec.max_target,
            "Achieved": achieved,
            "Status": f"Binding ({' and '.join(binding)})" if binding else "Slack",
            "Slack to target": max(min(distances), 0.0),
        })
    return pd.DataFrame(records)


def cost_table(data: BlendData, result: SolveResult) -> pd.DataFrame:
    """Blend cost per dry tonne, per as-received tonne, and per batch."""
    _require_solution(result)
    cost_dry = result.cost_per_dry_t
    as_received = as_received_tonnes_per_dry_tonne(data, result.fractions)
    dry_kg = batch_dry_kg(data, result.fractions)
    return pd.DataFrame([
        {"Item": "Cost (Rs per dry tonne of blend)", "Value": cost_dry, "Unit": "Rs/dry t"},
        # 1 dry t of blend weighs W t as received, so each as-received t costs cost_dry / W
        {"Item": "Cost (Rs per as-received tonne of blend)", "Value": cost_dry / as_received,
         "Unit": "Rs/as-received t"},
        {"Item": "Batch size (kg, as received)", "Value": data.settings.batch_size_kg, "Unit": "kg"},
        {"Item": "Batch dry mass (kg)", "Value": dry_kg, "Unit": "kg"},
        {"Item": "Cost per batch (Rs)", "Value": cost_dry * dry_kg / 1000.0, "Unit": "Rs"},
    ])


def summary_table(data: BlendData, result: SolveResult) -> pd.DataFrame:
    """Run information (text only): product, status, bases, input files, warnings."""
    items = [
        ("Product", data.settings.product_name),
        ("Solver status", result.status),
        ("Spec basis", data.settings.spec_basis),
        ("Lab analysis convention", data.lab_convention),
        ("Batch size", f"{_fmt_kg(data, data.settings.batch_size_kg)} kg as received"),
        ("Scale resolution", f"{data.settings.scale_resolution_kg:g} kg"),
        ("Raw materials read from", data.sources.get(config.SHEET_RAW_MATERIALS, "")),
        ("Specs read from", data.sources.get(config.SHEET_SPECS, "")),
        ("Input warnings", str(len(data.warnings))),
    ]
    items += [(f"Warning {number}", text) for number, text in enumerate(data.warnings, start=1)]
    return pd.DataFrame(items, columns=["Item", "Value"])


# ===========================================================================
# 5. Sensitivity: shadow prices and reduced costs
# ===========================================================================
#
# Sign conventions (checked in Stage 0 against finite differences, and in the
# Stage 3 tests through the LP duality conditions):
#
#   * PuLP + HiGHS give, for every row, pi = d(cost)/d(rhs): how much the
#     minimum cost rises per unit increase of the row's right-hand side.
#     In a minimisation a binding >= row has pi >= 0 and a binding <= row
#     has pi <= 0.
#   * dj (reduced cost) = cost_i - sum_k pi_k * a_ki, in Rs per dry tonne of
#     material per unit of fraction x_i (x_i = 1 means 100 % of the blend).
#
# What the report shows is "Rs/t saved per 1 unit of relaxation", positive
# when relaxing saves money:
#
#   * Relaxing a Min means LOWERING it; relaxing a Max means RAISING it.
#     So  saving = +d(cost)/dL  for a Min, and  -d(cost)/dL  for a Max.
#
#   * d(cost)/dL from pi. A dry composition row is  sum(N_i x_i) >= L,
#     with L on the right-hand side, so d(cost)/dL = pi.
#     Every other row is  sum((N_i - L g_i) x_i) >= 0:  L sits inside the
#     coefficients and the right-hand side is 0. Raising L by dL lowers the
#     left-hand side by sum(g_i x_i) * dL for the current recipe, which has
#     the same effect as raising the right-hand side by that amount. So
#           d(cost)/dL = pi * sum(g_i x_i).
#     On a dry basis g_i = 1 and sum(x_i) = 1, so this one formula covers
#     both kinds of row.
#
#   * These are derivatives: they hold only for small changes. Further
#     away the recipe changes, and so does the saving. On a dry composition
#     row the cost changes linearly while the same limits stay binding. On
#     other rows L sits inside the coefficients, so the cost curves a little
#     even then: the shadow price is exact at the current limit only.

SHADOW_PRICE_COLUMNS = [
    "Constraint", "Type", "Limits", "Customer limit", "Internal target", "Limit unit", "Relax by",
    "Saving per 1 unit (Rs/dry t)", "Step", "Saving per step (Rs/dry t)",
]
REDUCED_COST_COLUMNS = [
    "Raw material", "Price as entered (Rs/t)", "Price basis", "Reduced cost (Rs/dry t)",
    "Price drop needed (Rs/t as entered)", "Enters below (Rs/t as entered)", "Note",
]


def relaxation_step(limit: float) -> float:
    """A sensible step to quote a shadow price for: 1 unit, or less for small limits.

    For a limit of 1 or more the step is 1 (e.g. per 1 % CaO). Below 1 it is
    one unit of the limit's first significant decimal place: 0.03 % P -> 0.01,
    0.5 -> 0.1. This keeps "Rs per step" readable for trace elements.
    """
    if limit == 0 or abs(limit) >= 1:
        return 1.0
    return 10.0 ** math.floor(math.log10(abs(limit)))


def _limit_unit(data: BlendData, row: LinearRow) -> str:
    return "ratio" if row.spec.is_ratio else f"wt % ({data.settings.spec_basis} basis)"


def limit_saving_per_unit(row: LinearRow, result: SolveResult) -> float:
    """Rs per dry tonne of blend saved by relaxing ``row``'s limit by one unit.

    See the notes above this function for the derivation and sign rules.
    """
    weight = sum(row.denominator[name] * x for name, x in result.fractions.items())  # sum(g_i x_i)
    d_cost_d_limit = result.row_dual[row.name] * weight
    return d_cost_d_limit if row.bound == MIN else -d_cost_d_limit


def shadow_price_table(data: BlendData, result: SolveResult) -> pd.DataFrame:
    """Shadow prices of the binding limits: spec rows, then usage limits.

    Usage limits (variable bounds) have no row of their own; their shadow
    price is the material's reduced cost dj. dj is per unit fraction, i.e.
    per 100 % of the blend, so per 1 % it is dj / 100. A material held at its
    Max has dj <= 0 and raising the Max saves -dj/100 per %; one held at a
    Min above 0 has dj >= 0 and lowering the Min saves dj/100 per %.
    """
    _require_solution(result)
    records = []
    for row in result.rows.values():
        if result.row_slack[row.name] >= config.BINDING_TOLERANCE:
            continue
        saving = limit_saving_per_unit(row, result)
        step = relaxation_step(row.limit)
        records.append({
            "Constraint": row.name,
            "Type": "Ratio" if row.spec.is_ratio else "Composition",
            "Limits": row.spec.expression,
            "Customer limit": row.customer_limit,
            "Internal target": row.limit,
            "Limit unit": _limit_unit(data, row),
            "Relax by": "lowering the Min" if row.bound == MIN else "raising the Max",
            "Saving per 1 unit (Rs/dry t)": saving,
            "Step": step,
            "Saving per step (Rs/dry t)": saving * step,
        })

    for name, x in result.fractions.items():
        limit = data.effective_limit(name)
        dj = result.reduced_costs[name]
        pct = 100.0 * x
        at_upper = limit.max_pct < 100.0 and pct >= limit.max_pct - 1e-7
        at_lower = pct <= limit.min_pct + 1e-7
        # A material can sit at both bounds (e.g. banned: Min 0, Max 0). Then the sign
        # of dj says which one holds it: dj < 0 = it wants more (the Max holds it),
        # dj > 0 = it wants less (the Min holds it, or it is simply not worth using).
        if at_upper and (not at_lower or dj < 0):
            bound, relax, saving, value = "max", "raising the Max", -dj / 100.0, limit.max_pct
        elif at_lower and limit.min_pct > 0.0 and (not at_upper or dj > 0):
            bound, relax, saving, value = "min", "lowering the Min", dj / 100.0, limit.min_pct
        else:
            continue  # between its limits, or unused at 0 % (see reduced_cost_table)
        records.append({
            "Constraint": f"{name} {bound} use",
            "Type": "Usage limit",
            "Limits": name,
            "Customer limit": None,
            "Internal target": value,
            "Limit unit": "wt % of dry blend",
            "Relax by": relax,
            "Saving per 1 unit (Rs/dry t)": saving,
            "Step": 1.0,
            "Saving per step (Rs/dry t)": saving,
        })
    return pd.DataFrame(records, columns=SHADOW_PRICE_COLUMNS)


def reduced_cost_table(data: BlendData, result: SolveResult) -> pd.DataFrame:
    """Unused raw materials: how much cheaper each must become to enter the recipe.

    The reduced cost dj (Rs per dry tonne) is exactly that price drop: with
    a dry cost lower by more than dj, using some of the material would lower
    the blend cost. The drop is also shown in the user's own price basis:
    dry cost = price / (1 - m/100), so a dry drop of dj is a drop of
    dj * (1 - m/100) in an as-received price.

    Only materials with no minimum use and none in the recipe are listed.
    """
    _require_solution(result)
    records = []
    for name, x in result.fractions.items():
        limit = data.effective_limit(name)
        if x > _ZERO or limit.min_pct > 0:
            continue
        material = data.raw_materials[name]
        dj = result.reduced_costs[name]
        drop = price_from_dry_basis(dj, material.price_basis, material.moisture_pct)
        enters_below = material.cost_input - drop
        if limit.max_pct == 0:
            note = ("Banned by its usage limit (Max 0 %); it would be used if allowed"
                    if dj < 0 else "Banned by its usage limit (Max 0 %)")
        elif dj == 0:
            note = "Break-even: another recipe of the same cost uses it"
        elif enters_below <= 0:
            note = "Would not be used even if free"
        else:
            note = ""
        records.append({
            "Raw material": name,
            "Price as entered (Rs/t)": material.cost_input,
            "Price basis": material.price_basis,
            "Reduced cost (Rs/dry t)": dj,
            "Price drop needed (Rs/t as entered)": drop,
            "Enters below (Rs/t as entered)": enters_below,
            "Note": note,
        })
    return pd.DataFrame(records, columns=REDUCED_COST_COLUMNS)


def build_report(data: BlendData, result: SolveResult) -> dict[str, pd.DataFrame]:
    """All report tables, sheet name -> table. Only the summary if there is no solution."""
    tables = {SHEET_SUMMARY: summary_table(data, result)}
    if result.is_optimal:
        tables[SHEET_RECIPE] = recipe_table(data, result)
        tables[SHEET_COMPOSITION] = composition_table(data, result)
        tables[SHEET_SPEC_CHECK] = spec_check_table(data, result)
        tables[SHEET_COST] = cost_table(data, result)
        tables[SHEET_SHADOW_PRICES] = shadow_price_table(data, result)
        tables[SHEET_REDUCED_COSTS] = reduced_cost_table(data, result)
    return tables


def report_notes(data: BlendData, result: SolveResult) -> dict[str, list[str]]:
    """Explanatory notes written under each table in the Excel report."""
    notes = {
        SHEET_RECIPE: [
            f"Batch = {_fmt_kg(data, data.settings.batch_size_kg)} kg AS RECEIVED "
            f"(what is weighed into the mixer).",
        ],
        SHEET_COMPOSITION: [
            f"Components are wt % on the spec basis ({data.settings.spec_basis}). "
            f"LOI is shown on a dry basis and moisture on an as-received basis.",
        ],
        SHEET_SPEC_CHECK: [
            "Target = customer limit with the margin applied inward (Min + margin, Max - margin).",
            "Binding = the target is holding the recipe where it is; Slack = it has room to spare.",
            "Slack to target is in the spec's own unit.",
        ],
        SHEET_SHADOW_PRICES: [
            "Rs per dry tonne of blend saved by relaxing the limit (lowering a Min, raising a Max); "
            "tightening it costs the same per unit.",
            "Valid only near the current limit: further away the recipe changes, and so does the saving.",
            "For ratio limits and limits on an ignited or as-received basis the saving also drifts "
            "slightly as the limit moves, even before the recipe changes.",
            "'Saving per step' is the same figure for a step sized to the limit (e.g. per 0.01 % for P).",
        ],
        SHEET_REDUCED_COSTS: [
            "How much cheaper each unused raw material must become before it enters the recipe.",
        ],
    }
    if result.is_optimal:
        _, adjustments = weigh_sheet(recipe_table(data, result), data.settings.scale_resolution_kg)
        notes[SHEET_RECIPE] += weigh_sheet_notes(adjustments, data.settings.scale_resolution_kg)
        if shadow_price_table(data, result).empty:
            notes[SHEET_SHADOW_PRICES].insert(0, "No binding limits.")
        if reduced_cost_table(data, result).empty:
            notes[SHEET_REDUCED_COSTS].insert(0, "None: every raw material is used.")
    return notes


# ===========================================================================
# 6. Display: rounding, Excel workbook, console summary
# ===========================================================================
# How each number is shown:
#   * The kind of a column comes from its header: "kg" -> KG, "Rs" -> RS,
#     "Step" as it is, everything else PCT (wt %, ratios, limits). In a table
#     with a "Unit" column (the cost tables) the Unit decides, row by row.
#   * PCT numbers in a row that holds limits typed by the user (columns
#     starting with "Customer", and "Margin") get one more decimal than the
#     most precise of those limits, and at least DISPLAY_DECIMALS_PCT.
#     So P Max 0.03 shows its target and achieved value as 0.030.
#   * A recipe table (a "Raw material" column with a TOTAL row) is rounded
#     as a weigh sheet: see weigh_sheet().

PCT, KG, RS, PLAIN = "pct", "kg", "rs", "plain"
_KIND_OVERRIDES = {"Step": PLAIN}
_TOTAL_LABEL = "TOTAL"


def column_kind(header: str) -> str:
    """Display kind of a column, from its header."""
    if header in _KIND_OVERRIDES:
        return _KIND_OVERRIDES[header]
    if "kg" in header:
        return KG
    if "Rs" in header:
        return RS
    return PCT


def _unit_kind(unit: object) -> str:
    """Display kind for a row of a table with a Unit column."""
    text = str(unit)
    return KG if "kg" in text else RS if "Rs" in text else PCT


def _is_number(value: object) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def _is_missing(value: object) -> bool:
    return value is None or (isinstance(value, float) and math.isnan(value))


def decimals_entered(value: float) -> int:
    """Decimals of a number as typed: 50 -> 0, 0.03 -> 2, 0.125 -> 3.

    ``repr`` gives the shortest text that reads back as the same number,
    which is what was typed in Excel. Capped at DISPLAY_DECIMALS_PCT_MAX so
    a computed cell such as 0.30000000000000004 does not show 17 decimals.
    """
    exponent = Decimal(repr(float(value))).normalize().as_tuple().exponent
    return min(max(0, -exponent), config.DISPLAY_DECIMALS_PCT_MAX)


def limit_decimals(entered: Iterable[object]) -> int:
    """Decimals for a spec row: one more than its most precise typed limit, at least the default."""
    numbers = [value for value in entered if _is_number(value) and not _is_missing(value)]
    if not numbers:
        return config.DISPLAY_DECIMALS_PCT
    wanted = max(decimals_entered(value) for value in numbers) + 1
    return min(max(config.DISPLAY_DECIMALS_PCT, wanted), config.DISPLAY_DECIMALS_PCT_MAX)


# Columns (besides those starting with "Customer") that hold values typed by the
# user, e.g. the old and new limits in a what-if Changes table
LIMIT_SOURCE_COLUMNS = {"Margin", "Base value", "Scenario value"}


def _is_limit_source(column: str) -> bool:
    """Columns holding limits as typed by the user; they set their row's decimals."""
    return column.startswith("Customer") or column in LIMIT_SOURCE_COLUMNS


@dataclass(frozen=True)
class CellFormat:
    """How one cell is shown: its kind and number of decimals."""

    kind: str
    decimals: int


def composition_decimals(data: BlendData) -> list[int | None]:
    """Per-row PCT decimals for the composition table.

    A component limited on its own by a composition spec row (Numerator is
    just that component, no Denominator) follows the same rule as that row
    in the spec check: one more decimal than its typed limits, at least the
    default. So P shows 0.030 in both tables. Other rows (unlimited
    components, the sum, LOI, moisture) return None, meaning "use the
    default".
    """
    decimals: list[int | None] = []
    for component in data.components:
        # "On its own" = the whole numerator is that one component with factor 1
        own_specs = [spec for spec in data.specs if not spec.is_ratio and spec.numerator == (Term(1.0, component),)]
        decimals.append(max(limit_decimals([spec.min_value, spec.max_value, spec.margin]) for spec in own_specs)
                        if own_specs else None)
    return decimals + [None, None, None]  # Sum of components, LOI, Moisture


def report_row_decimals(data: BlendData, result: SolveResult) -> dict[str, list[int | None]]:
    """Per-row decimal overrides for report sheets whose rows hold no typed limits."""
    return {SHEET_COMPOSITION: composition_decimals(data)} if result.is_optimal else {}


def cell_formats(table: pd.DataFrame, row_decimals: Sequence[int | None] | None = None) -> list[list[CellFormat]]:
    """Display format of every cell, row by row (see the rules above).

    Args:
        table: The table.
        row_decimals: Optional PCT decimals per row (None = work it out from
            the row's typed-limit columns, as usual).
    """
    has_unit = "Unit" in table.columns
    sources = [column for column in table.columns if _is_limit_source(column)]
    formats = []
    for row_index, (_, row) in enumerate(table.iterrows()):
        pct_decimals = limit_decimals(row[column] for column in sources)
        if row_decimals is not None and row_decimals[row_index] is not None:
            pct_decimals = row_decimals[row_index]
        row_formats = []
        for column in table.columns:
            kind = _unit_kind(row["Unit"]) if has_unit and column != "Unit" else column_kind(column)
            decimals = {KG: config.DISPLAY_DECIMALS_KG, RS: config.DISPLAY_DECIMALS_RS}.get(kind, pct_decimals)
            row_formats.append(CellFormat(kind, decimals))
        formats.append(row_formats)
    return formats


def round_to_scale(value: float, resolution: float = config.SCALE_RESOLUTION_KG) -> float:
    """Round a mass to the nearest multiple of the scale resolution (e.g. 0.5 kg)."""
    return round(round(value / resolution) * resolution, 10)


def format_number(
    value: object,
    kind: str,
    scale_resolution_kg: float = config.SCALE_RESOLUTION_KG,
    decimals: int | None = None,
) -> str:
    """Text for one cell: rounded for display, blank for missing values.

    Args:
        value: The cell value (numbers are formatted; text passes through).
        kind: PCT, KG, RS or PLAIN.
        scale_resolution_kg: Resolution kg values are rounded to.
        decimals: Override the default number of decimals for the kind.
    """
    if _is_missing(value):
        return ""
    if not _is_number(value):
        return str(value)
    if math.isinf(value):
        return "inf" if value > 0 else "-inf"
    if kind == PLAIN:
        return f"{value:g}"
    if kind == KG:
        value = round_to_scale(value, scale_resolution_kg)
    if decimals is None:
        decimals = {KG: config.DISPLAY_DECIMALS_KG, RS: config.DISPLAY_DECIMALS_RS}.get(kind, config.DISPLAY_DECIMALS_PCT)
    if round(value, decimals) == 0:
        value = 0.0  # avoid showing "-0.00"
    separator = "," if kind in (KG, RS) else ""
    return f"{value:{separator}.{decimals}f}"


# ---------------------------------------------------------------------------
# Weigh sheet
# ---------------------------------------------------------------------------

def weigh_sheet_rounding(values: Sequence[float], resolution: float) -> tuple[list[float], float, int | None, float]:
    """Round masses to the scale so that they add up exactly to the rounded total.

    Each value and the total are rounded to the nearest multiple of
    ``resolution``. The rounded values need not add up to the rounded total
    (each can be off by up to half a step), so the difference is put on the
    largest value, where it matters least in relative terms.

    The arithmetic is done in whole scale steps, so no float error creeps in.

    Args:
        values: Masses in kg.
        resolution: Scale resolution in kg.

    Returns:
        (rounded values, rounded total, index of the adjusted value or None,
        difference added to it in kg).
    """
    steps = [round(value / resolution) for value in values]
    total_steps = round(sum(values) / resolution)
    difference = total_steps - sum(steps)
    adjusted = None
    if difference and values:
        adjusted = max(range(len(values)), key=lambda index: values[index])
        steps[adjusted] += difference
    rounded = [round(step * resolution, 10) for step in steps]
    return rounded, round(total_steps * resolution, 10), adjusted, round(difference * resolution, 10)


@dataclass(frozen=True)
class WeighAdjustment:
    """Where the rounding difference of one kg column was put."""

    column: str
    material: str | None     # None if no adjustment was needed
    difference_kg: float


def _is_weigh_sheet(table: pd.DataFrame) -> bool:
    return "Raw material" in table.columns and bool((table["Raw material"] == _TOTAL_LABEL).any())


def weigh_sheet(table: pd.DataFrame, resolution: float) -> tuple[pd.DataFrame, list[WeighAdjustment]]:
    """Copy of a recipe table with every kg column rounded as a weigh sheet.

    In each kg column the material rows are rounded to the scale and add up
    exactly to the rounded TOTAL row; see :func:`weigh_sheet_rounding`.
    Columns with a missing value (e.g. an infeasible scenario) are left alone.

    Returns:
        (rounded table, one WeighAdjustment per kg column).
    """
    table = table.copy()
    is_total = table["Raw material"] == _TOTAL_LABEL
    adjustments = []
    for column in table.columns:
        if column_kind(column) != KG:
            continue
        values = [float(value) for value in table.loc[~is_total, column]]
        if any(math.isnan(value) for value in values):
            continue
        rounded, total, index, difference = weigh_sheet_rounding(values, resolution)
        table.loc[~is_total, column] = rounded
        table.loc[is_total, column] = total
        material = None if index is None else table.loc[~is_total, "Raw material"].iloc[index]
        adjustments.append(WeighAdjustment(column, material, difference))
    return table, adjustments


def weigh_sheet_notes(adjustments: list[WeighAdjustment], resolution: float) -> list[str]:
    """Report lines saying which row took the rounding difference in each kg column."""
    lines = [f"kg are rounded to the scale resolution ({resolution:g} kg); each kg column adds up "
             f"exactly to its TOTAL."]
    for adjustment in adjustments:
        if adjustment.material is None:
            lines.append(f"{adjustment.column}: no rounding adjustment needed.")
        else:
            lines.append(
                f"{adjustment.column}: rounding difference of "
                f"{adjustment.difference_kg:+.{config.DISPLAY_DECIMALS_KG}f} kg put on "
                f"{adjustment.material} (the largest)."
            )
    return lines


def display_table(
    table: pd.DataFrame,
    scale_resolution_kg: float = config.SCALE_RESOLUTION_KG,
    row_decimals: Sequence[int | None] | None = None,
) -> pd.DataFrame:
    """A copy of ``table`` with every cell as display text (rounded; blanks for missing).

    Args:
        table: The table.
        scale_resolution_kg: Resolution kg values are rounded to.
        row_decimals: Optional PCT decimals per row, e.g. from
            :func:`composition_decimals`.
    """
    if _is_weigh_sheet(table):
        table, _ = weigh_sheet(table, scale_resolution_kg)
    formats = cell_formats(table, row_decimals)
    shown = [
        [format_number(value, cell.kind, scale_resolution_kg, cell.decimals)
         for value, cell in zip(row, row_formats)]
        for row, row_formats in zip(table.itertuples(index=False), formats)
    ]
    return pd.DataFrame(shown, columns=table.columns)


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------

def _excel_format(cell: CellFormat) -> str:
    """Excel number format for a cell (rounding on screen only)."""
    if cell.kind == PLAIN:
        return "General"
    grouping = "#,##0" if cell.kind in (KG, RS) else "0"
    return grouping + ("." + "0" * cell.decimals if cell.decimals else "")


@dataclass
class Section:
    """One titled table on a sheet that holds several (e.g. a what-if scenario sheet).

    Attributes:
        title: Heading written in bold above the table.
        table: The table (full precision, like any report table).
        notes: Lines written in italics under the table.
    """

    title: str
    table: pd.DataFrame
    notes: list[str] = field(default_factory=list)


def write_excel(
    tables: dict[str, pd.DataFrame | list[Section]],
    path: str | Path,
    notes: dict[str, list[str]] | None = None,
    scale_resolution_kg: float = config.SCALE_RESOLUTION_KG,
    row_decimals: dict[str, list[int | None]] | None = None,
) -> Path:
    """Write the report tables to an Excel workbook.

    A sheet holds either one table, or a list of :class:`Section` s stacked
    one under another (title, table, notes, blank row).

    Cells keep full precision; Excel number formats do the rounding on
    screen, so the values can still be used in further calculations. The
    exception is kg: they are stored rounded to the scale resolution (and,
    in a recipe table, as a weigh sheet that adds up to its total), because
    a number format cannot round to e.g. 0.5 kg. Missing values are blank
    cells. Notes are written in italics below each table.

    Args:
        tables: Sheet name -> table or sections (e.g. from :func:`build_report`).
        path: Output .xlsx path (folders are created).
        notes: Sheet name -> lines of explanation (for single-table sheets).
        scale_resolution_kg: Resolution kg values are rounded to
            (normally ``data.settings.scale_resolution_kg``).
        row_decimals: Sheet name -> PCT decimals per row (see
            :func:`report_row_decimals`).

    Returns:
        The path written.
    """
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    notes = notes or {}
    row_decimals = row_decimals or {}
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for sheet_name, content in tables.items():
            if isinstance(content, pd.DataFrame):
                widths = _write_table(writer, sheet_name, content, 1, scale_resolution_kg,
                                      row_decimals.get(sheet_name), freeze=True)
                _write_notes(writer.sheets[sheet_name], notes.get(sheet_name, []), first_row=len(content) + 3)
            else:
                widths = {}
                row = 1
                for section in content:
                    # title in `row`, header in row + 1, data below it, then the notes
                    section_widths = _write_table(writer, sheet_name, section.table, row + 1,
                                                  scale_resolution_kg, None, freeze=False)
                    for column, width in section_widths.items():
                        widths[column] = max(widths.get(column, 0), width)
                    worksheet = writer.sheets[sheet_name]
                    worksheet.cell(row=row, column=1, value=section.title).font = Font(bold=True, size=12)
                    notes_row = row + len(section.table) + 3
                    _write_notes(worksheet, section.notes, first_row=notes_row)
                    row = notes_row + len(section.notes) + 1
            worksheet = writer.sheets[sheet_name]
            for column_number, width in widths.items():
                worksheet.column_dimensions[worksheet.cell(row=1, column=column_number).column_letter].width = width
    return path


def _write_table(
    writer: pd.ExcelWriter,
    sheet_name: str,
    table: pd.DataFrame,
    header_row: int,
    scale_resolution_kg: float,
    row_decimals: Sequence[int | None] | None,
    freeze: bool,
) -> dict[int, float]:
    """Write one table with its header in ``header_row``; return the column widths it needs."""
    if _is_weigh_sheet(table):
        table, _ = weigh_sheet(table, scale_resolution_kg)
    formats = cell_formats(table, row_decimals)
    stored = table.copy().astype(object)
    for row_number, row_formats in enumerate(formats):
        for column_number, cell_format in enumerate(row_formats):
            value = stored.iat[row_number, column_number]
            if cell_format.kind == KG and _is_number(value) and not math.isnan(value):
                stored.iat[row_number, column_number] = round_to_scale(value, scale_resolution_kg)
    # NaN/None become blank cells. Writing again to the same sheet adds to it.
    stored.to_excel(writer, sheet_name=sheet_name, index=False, startrow=header_row - 1)

    worksheet = writer.sheets[sheet_name]
    style_header_row(worksheet, header_row, freeze=freeze)
    for row_number, row_formats in enumerate(formats, start=header_row + 1):
        for column_number, cell_format in enumerate(row_formats, start=1):
            cell = worksheet.cell(row=row_number, column=column_number)
            if _is_number(cell.value):
                cell.number_format = _excel_format(cell_format)

    # Width: header (wrapped at 24 characters) or the longest displayed value
    shown = display_table(table, scale_resolution_kg, row_decimals)
    return {
        column_number: min(max([min(len(str(column)), 24)] + [len(text) for text in shown[column]]) + 2, 70)
        for column_number, column in enumerate(shown.columns, start=1)
    }


def _write_notes(worksheet: Worksheet, lines: list[str], first_row: int) -> None:
    for offset, line in enumerate(lines):
        cell = worksheet.cell(row=first_row + offset, column=1, value=line)
        cell.font = Font(italic=True)


# ---------------------------------------------------------------------------
# Console
# ---------------------------------------------------------------------------

def _fmt_kg(data: BlendData, value: float) -> str:
    return format_number(value, KG, data.settings.scale_resolution_kg)


def _fmt_rs(value: float) -> str:
    return format_number(value, RS)


def indented_table(table: pd.DataFrame, columns: list[str],
                   scale_resolution_kg: float = config.SCALE_RESOLUTION_KG) -> str:
    """Selected columns of a table as display text, indented for the console.

    The whole table is formatted first, so the per-row decimals (which look
    at the limit columns) are right even when those columns are not shown.
    """
    text = display_table(table, scale_resolution_kg)[columns].to_string(index=False)
    return "\n".join("  " + line for line in text.splitlines())


def format_summary(data: BlendData, result: SolveResult) -> str:
    """The console summary as text (see :func:`print_summary`)."""
    settings = data.settings
    resolution = settings.scale_resolution_kg
    lines = [
        f"=== Least-cost blend: {settings.product_name} ===",
        f"Status: {result.status} | spec basis: {settings.spec_basis} | "
        f"lab convention: {data.lab_convention}",
    ]
    if data.warnings:
        lines.append(f"Input warnings ({len(data.warnings)}):")
        lines += [f"  - {warning}" for warning in data.warnings]
    if not result.is_optimal:
        lines.append(f"No recipe: the solver status is '{result.status}'.")
        return "\n".join(lines)

    cost = cost_table(data, result).set_index("Item")["Value"]
    recipe = recipe_table(data, result)
    _, adjustments = weigh_sheet(recipe, resolution)
    lines += [
        f"Cost: Rs {_fmt_rs(result.cost_per_dry_t)} per dry t of blend | "
        f"Rs {_fmt_rs(cost['Cost (Rs per as-received tonne of blend)'])} per as-received t | "
        f"Rs {_fmt_rs(cost['Cost per batch (Rs)'])} per batch",
        "",
        f"Recipe (batch {_fmt_kg(data, settings.batch_size_kg)} kg as received, "
        f"{_fmt_kg(data, cost['Batch dry mass (kg)'])} kg dry):",
        indented_table(recipe, ["Raw material", "Dry blend (wt %)", "Dry mass (kg/batch)",
                                "As-received mass (kg/batch)"], resolution),
    ]
    lines += ["  " + line for line in weigh_sheet_notes(adjustments, resolution)]
    lines += [
        "",
        "Spec check:",
        indented_table(spec_check_table(data, result),
                       ["Spec", "Limits", "Basis", "Target Min", "Target Max", "Achieved", "Status"], resolution),
        "",
    ]
    shadow = shadow_price_table(data, result)
    lines.append("Shadow prices of binding limits (Rs per dry t saved by relaxing; valid only near the "
                 "current limit):")
    if shadow.empty:
        lines.append("  none")
    else:
        lines.append(indented_table(shadow, ["Constraint", "Relax by", "Step", "Limit unit",
                                             "Saving per step (Rs/dry t)", "Saving per 1 unit (Rs/dry t)"]))
    reduced = reduced_cost_table(data, result)
    lines += ["", "Unused raw materials (price drop needed before they enter):"]
    if reduced.empty:
        lines.append("  none: every raw material is used")
    else:
        lines.append(indented_table(reduced, ["Raw material", "Reduced cost (Rs/dry t)",
                                              "Price drop needed (Rs/t as entered)",
                                              "Enters below (Rs/t as entered)", "Note"]))
    if result.lp_file:
        lines += ["", f"LP file: {result.lp_file}"]
    return "\n".join(lines)


def print_summary(data: BlendData, result: SolveResult) -> None:
    """Print a readable summary of the result to the console."""
    print(format_summary(data, result))

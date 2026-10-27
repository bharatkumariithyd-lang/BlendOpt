"""Write input workbooks: blank templates, or any tables you pass in.

This module is generic: it knows the sheet and column layout (from
``config``) but no raw material, component or product names. The example
datasets live in ``examples.py``.
"""

from pathlib import Path
from typing import Sequence

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from blend_optimizer import config

_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill("solid", fgColor="305496")
_SECTION_FONT = Font(bold=True, size=12)
_LAST_VALIDATED_ROW = 500  # drop-down lists are added down to this row


# ---------------------------------------------------------------------------
# README sheet text: (column A, column B). Rows with empty column B are headings.
# ---------------------------------------------------------------------------
def _readme_rows() -> list[tuple[str, str]]:
    """Instructions written to the README sheet of every generated workbook."""
    low, high = config.ANALYSIS_TOTAL_WARN_RANGE
    convention_text = {
        config.LAB_DRY_INCL_LOI: (
            f"wt %, dry basis, with LOI reported separately, so components + LOI should "
            f"total {low:g}-{high:g} %."
        ),
        config.LAB_LOI_FREE: (
            f"wt % on an ignited (LOI-free) basis, so the components alone should total "
            f"{low:g}-{high:g} %. They are converted to dry basis using LOI_%."
        ),
    }[config.LAB_ANALYSIS_CONVENTION]
    return [
        ("Least-cost blend optimizer: input workbook", ""),
        ("", ""),
        ("Layouts", ""),
        ("1. One workbook per product", f"Sheets {config.SHEET_RAW_MATERIALS}, {config.SHEET_SPECS}, "
                                        f"{config.SHEET_SETTINGS} and (optional) {config.SHEET_USAGE_LIMITS}."),
        ("2. Master + product", f"A master workbook with only {config.SHEET_RAW_MATERIALS}, plus a product "
                                f"workbook with {config.SHEET_SPECS}, {config.SHEET_SETTINGS} and "
                                f"(optional) {config.SHEET_USAGE_LIMITS}."),
        ("", ""),
        ("General rules", ""),
        ("Headers", "Keep the column headers in row 1 and do not rename the fixed columns. "
                    "Blank rows are ignored."),
        ("Numbers", "Type plain numbers (no units, no '<0.01')."),
        ("Blank or 0", "In a component column, blank = not analysed, 0 = analysed, none present. Both count "
                       "as 0, but a blank in a component that a spec uses gives a warning: type 0 if it is "
                       "genuinely absent."),
        ("Analyses", convention_text + " (Set by LAB_ANALYSIS_CONVENTION in config.py.)"),
        ("", ""),
        (f"Sheet {config.SHEET_RAW_MATERIALS}", ""),
        (config.COL_NAME, "Unique raw material name."),
        (config.COL_COST, f"Price in Rs per tonne, on the basis given in {config.COL_PRICE_BASIS}."),
        (config.COL_PRICE_BASIS, f"'{config.PRICE_AS_RECEIVED}' or '{config.PRICE_DRY}'. Prices are converted "
                                 f"to Rs per dry tonne: dry cost = cost / (1 - moisture/100) for "
                                 f"{config.PRICE_AS_RECEIVED}. May be blank only if moisture is 0."),
        (config.COL_MOISTURE, "Free moisture, wt % as received. Blank = 0."),
        (config.COL_LOI, "Loss on ignition, wt % dry basis. Blank = 0."),
        (config.COL_MIN_USE, "Plant-wide minimum use, wt % of the dry blend. Blank = 0."),
        (config.COL_MAX_USE, "Plant-wide maximum use, wt % of the dry blend. Blank = 100."),
        ("Any other column", "A chemical component, wt % (see 'Analyses' above). Add, remove or "
                             "reorder component columns freely; the header is the component name."),
        ("", ""),
        (f"Sheet {config.SHEET_SPECS}", "One row per limit."),
        (config.COL_NAME, "Unique label. The constraints are named <Name>_min and <Name>_max in the report."),
        (config.COL_NUMERATOR, f"A component, or components joined with '{config.COMPONENT_SEPARATOR}' "
                               f"(e.g. CaO{config.COMPONENT_SEPARATOR}MgO). A component may be counted as another "
                               f"by a factor written before it with '{config.FACTOR_SEPARATOR}', e.g. "
                               f"Al2O3{config.COMPONENT_SEPARATOR}1.8895{config.FACTOR_SEPARATOR}Al counts "
                               f"metallic Al as Al2O3 (1.8895 = 101.96 / (2 x 26.98)). Factors are numbers you "
                               f"type; '-' is not allowed. Names must match the {config.SHEET_RAW_MATERIALS} "
                               f"headers exactly."),
        (config.COL_DENOMINATOR, "Blank for a composition limit. Filled in (same syntax as Numerator), the row "
                                 "limits the ratio Numerator / Denominator."),
        (f"{config.COL_MIN}, {config.COL_MAX}", "The customer limits. Blank = no limit; at least one is needed."),
        (config.COL_MARGIN, "Safety margin in the same unit as Min/Max, applied inward "
                            "(target = Min + margin, Max - margin). Blank = 0."),
        ("", ""),
        (f"Sheet {config.SHEET_SETTINGS}", f"Two columns, {config.COL_KEY} and {config.COL_VALUE}."),
        (config.KEY_PRODUCT_NAME, "Shown in the report."),
        (config.KEY_SPEC_BASIS, f"Basis of the composition limits: {', '.join(config.SPEC_BASES)} "
                                f"(default {config.DEFAULT_SPEC_BASIS}). Ratio limits do not depend on it."),
        (config.KEY_BATCH_SIZE, f"Total as-received mass weighed into the mixer, kg "
                                f"(default {config.DEFAULT_BATCH_SIZE_KG:g})."),
        (config.KEY_SCALE_RESOLUTION, f"Optional. Resolution of the weighing scale, kg; the weigh sheet is "
                                      f"rounded to it. Blank = {config.SCALE_RESOLUTION_KG:g} (config.py)."),
        ("", ""),
        (f"Sheet {config.SHEET_USAGE_LIMITS} (optional)", "Product- or customer-specific usage limits."),
        (config.COL_NAME, f"Must match a name in {config.SHEET_RAW_MATERIALS}."),
        (f"{config.COL_MIN_USE}, {config.COL_MAX_USE}", "wt % of the dry blend. The tighter of these and the "
                                                        "plant-wide limits applies. Max 0 bans a material."),
        ("", ""),
        (f"Sheet {config.SHEET_WHAT_IF} (optional)", "What-if scenarios, each compared with the base case. "
                                                     "One row per change; rows with the same Scenario name "
                                                     "form one scenario."),
        (config.COL_SCENARIO, "Any name, e.g. 'Fluorspar dearer'."),
        (config.COL_CHANGE_TYPE, f"{config.CHANGE_PRICE} (Item = raw material, value = Rs/t in its own price "
                                 f"basis); {config.CHANGE_SPEC_MIN} / {config.CHANGE_SPEC_MAX} / "
                                 f"{config.CHANGE_SPEC_MARGIN} (Item = spec Name; blank Min/Max removes that "
                                 f"limit, blank Margin = 0); {config.CHANGE_USAGE_MIN} / {config.CHANGE_USAGE_MAX} "
                                 f"(Item = raw material, value = wt % of dry blend; replaces its UsageLimits "
                                 f"value, the plant-wide limit still applies)."),
        (config.COL_ITEM, "The raw material or spec Name being changed."),
        (config.COL_NEW_VALUE, "The new value (see Change type)."),
    ]


# ---------------------------------------------------------------------------
# Writing
# ---------------------------------------------------------------------------

def write_workbook(
    path: str | Path,
    raw_materials: pd.DataFrame | None = None,
    specs: pd.DataFrame | None = None,
    settings: dict[str, object] | None = None,
    usage_limits: pd.DataFrame | None = None,
    what_if: pd.DataFrame | None = None,
    include_readme: bool = True,
) -> Path:
    """Write an input workbook with the sheets that are given.

    Args:
        path: Output .xlsx path (folders are created if needed).
        raw_materials: Table for the RawMaterials sheet, or None to leave it out.
        specs: Table for the Specs sheet, or None.
        settings: Key -> value for the Settings sheet, or None.
        usage_limits: Table for the UsageLimits sheet, or None.
        what_if: Table for the WhatIf sheet, or None.
        include_readme: Add the README sheet with instructions (first sheet).

    Returns:
        The path written.
    """
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)

    sheets: dict[str, pd.DataFrame] = {}
    if raw_materials is not None:
        sheets[config.SHEET_RAW_MATERIALS] = raw_materials
    if specs is not None:
        sheets[config.SHEET_SPECS] = specs
    if settings is not None:
        sheets[config.SHEET_SETTINGS] = pd.DataFrame(
            list(settings.items()), columns=list(config.SETTINGS_COLUMNS)
        )
    if usage_limits is not None:
        sheets[config.SHEET_USAGE_LIMITS] = usage_limits
    if what_if is not None:
        sheets[config.SHEET_WHAT_IF] = what_if
    if not sheets:
        raise ValueError("write_workbook needs at least one sheet")

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        if include_readme:
            _write_readme(writer)
        for sheet_name, table in sheets.items():
            table.to_excel(writer, sheet_name=sheet_name, index=False)
            _format_table_sheet(writer.sheets[sheet_name])
        if config.SHEET_RAW_MATERIALS in sheets:
            _add_list_validation(writer.sheets[config.SHEET_RAW_MATERIALS], config.COL_PRICE_BASIS,
                                 config.PRICE_BASES)
        if config.SHEET_SETTINGS in sheets:
            _add_spec_basis_validation(writer.sheets[config.SHEET_SETTINGS])
        if config.SHEET_WHAT_IF in sheets:
            _add_list_validation(writer.sheets[config.SHEET_WHAT_IF], config.COL_CHANGE_TYPE, config.CHANGE_TYPES)
    return path


def blank_tables(components: Sequence[str] = ()) -> dict[str, object]:
    """Empty tables (headers only) for every sheet, plus default settings.

    Args:
        components: Component column headers to pre-fill in RawMaterials
            (they can be edited freely in Excel).

    Returns:
        Keyword arguments for :func:`write_workbook`.
    """
    return {
        "raw_materials": pd.DataFrame(columns=list(config.RAW_MATERIAL_FIXED_COLUMNS) + list(components)),
        "specs": pd.DataFrame(columns=list(config.SPEC_COLUMNS)),
        "settings": {
            config.KEY_PRODUCT_NAME: "",
            config.KEY_SPEC_BASIS: config.DEFAULT_SPEC_BASIS,
            config.KEY_BATCH_SIZE: config.DEFAULT_BATCH_SIZE_KG,
            config.KEY_SCALE_RESOLUTION: None,  # optional: blank uses config.py
        },
        "usage_limits": pd.DataFrame(columns=list(config.USAGE_LIMIT_COLUMNS)),
        "what_if": pd.DataFrame(columns=list(config.WHAT_IF_COLUMNS)),
    }


def write_blank_templates(folder: str | Path, components: Sequence[str] = ()) -> list[Path]:
    """Write the three blank templates (both layouts) into ``folder``.

    * ``Template_single_product.xlsx``: layout 1, all sheets.
    * ``Template_master_raw_materials.xlsx``: layout 2, RawMaterials only.
    * ``Template_product.xlsx``: layout 2, Specs + Settings + UsageLimits + WhatIf.

    Args:
        folder: Output folder (created if needed).
        components: Component headers to pre-fill in RawMaterials.

    Returns:
        The paths written.
    """
    folder = Path(folder)
    tables = blank_tables(components)
    product_only = {key: value for key, value in tables.items() if key != "raw_materials"}
    return [
        write_workbook(folder / "Template_single_product.xlsx", **tables),
        write_workbook(folder / "Template_master_raw_materials.xlsx", raw_materials=tables["raw_materials"]),
        write_workbook(folder / "Template_product.xlsx", **product_only),
    ]


# ---------------------------------------------------------------------------
# Formatting helpers
# ---------------------------------------------------------------------------

def _write_readme(writer: pd.ExcelWriter) -> None:
    """Add the README sheet as the first sheet."""
    worksheet = writer.book.create_sheet(config.SHEET_README, 0)
    writer.sheets[config.SHEET_README] = worksheet
    for row_number, (label, text) in enumerate(_readme_rows(), start=1):
        worksheet.cell(row=row_number, column=1, value=label or None)
        worksheet.cell(row=row_number, column=2, value=text or None)
        if label and not text:  # a heading
            worksheet.cell(row=row_number, column=1).font = _SECTION_FONT
        else:
            worksheet.cell(row=row_number, column=1).font = Font(bold=True)
        worksheet.cell(row=row_number, column=2).alignment = Alignment(wrap_text=True, vertical="top")
    worksheet.column_dimensions["A"].width = 34
    worksheet.column_dimensions["B"].width = 100


def style_header_row(worksheet: Worksheet, row: int = 1, freeze: bool = True) -> None:
    """Bold white-on-blue header cells in ``row``; row 1 is frozen so it stays visible.

    Also used by the report writer in ``solve_and_report`` (with other rows
    for sheets that hold several tables).
    """
    for cell in worksheet[row]:
        if cell.value is None:
            continue
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    if freeze:
        worksheet.freeze_panes = f"A{row + 1}"


def _format_table_sheet(worksheet: Worksheet) -> None:
    """Style the header row, freeze it, and size the columns to their contents."""
    style_header_row(worksheet)
    for column_cells in worksheet.columns:
        longest = max((len(str(cell.value)) for cell in column_cells if cell.value is not None), default=8)
        letter = get_column_letter(column_cells[0].column)
        worksheet.column_dimensions[letter].width = min(max(longest + 2, 9), 40)


def _add_list_validation(worksheet: Worksheet, header: str, choices: Sequence[str]) -> None:
    """Add a drop-down list to the column whose row-1 header is ``header``."""
    for cell in worksheet[1]:
        if cell.value == header:
            letter = get_column_letter(cell.column)
            validation = DataValidation(type="list", formula1='"' + ",".join(choices) + '"', allow_blank=True)
            validation.add(f"{letter}2:{letter}{_LAST_VALIDATED_ROW}")
            worksheet.add_data_validation(validation)
            return


def _add_spec_basis_validation(worksheet: Worksheet) -> None:
    """Add a drop-down list to the Value cell of the Spec_basis row in Settings."""
    for row in worksheet.iter_rows(min_row=2):
        if row[0].value == config.KEY_SPEC_BASIS:
            validation = DataValidation(type="list", formula1='"' + ",".join(config.SPEC_BASES) + '"',
                                        allow_blank=True)
            validation.add(row[1].coordinate)
            worksheet.add_data_validation(validation)
            return
